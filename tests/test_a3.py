"""Ponte: MIME real, valores assinados, tolerância e falhas independentes."""

from datetime import datetime
from decimal import Decimal
from email.message import EmailMessage
from io import BytesIO
import json
from pathlib import Path
import unittest

from openpyxl import Workbook, load_workbook

import test_aceite as amostras
from validar_saida import ErroAceite, verificar
from tarvo.execucao import executar
from tarvo.exportacao import publicar


class A3Test(unittest.TestCase):
    vendas = amostras.AceiteTest.vendas

    def setUp(self):
        amostras.AceiteTest.setUp(self)
        with self.regras.open("a", encoding="utf-8") as f:
            f.write('''
[remessas]
id="PONTE_TESTE"
participacao_fracao="0.5"
spread_padrao_fracao="0.015"
tolerancia_pp="0.01"
fonte="Amostra sintética"
aprovacao="Teste, sem aprovação financeira real"
aprovador="Teste"
aprovado_em="NAO_INFORMADO"
natureza="DECISAO_TIME"
''')
        livro = load_workbook(self.dados / "taxas.xlsx")
        livro.active.cell(1, 9, "Spread remessa")
        livro.active.cell(2, 9, .015)
        livro.save(self.dados / "taxas.xlsx")
        livro.close()
        self.email([{}])

    def email(self, linhas, nome="chegou_agora.eml", anexo="nome_livre.xlsx", layout=True, periodo=""):
        livro = Workbook()
        aba = livro.active
        aba.title = "Operações"
        aba.append(["Preâmbulo"])
        if periodo:
            aba.append([f"Competência: {periodo}"])
        colunas = ["Nº Operação", "Data Contratação", "Data Liquidação", "CNPJ Cliente", "Cliente", "Moeda",
                   "Valor ME", "Taxa Cliente", "PTAX Ref.", "Valor R$", "Spread R$", "Spread %",
                   "Spread Estimado R$ (D-1)", "IOF R$", "Status Compliance", "Observação"]
        aba.append(colunas if layout else ["Formato desconhecido"])
        for i, campos in enumerate(linhas, 1):
            linha = dict(zip(colunas, [f"PC-TESTE-{i}", datetime(2032, 4, 2), datetime(2032, 5, 1),
                "01234567800012", "Loja de teste", "USD", 100, 5.075, 5, 507.5, 262.01, .015, 900, 0, "APROVADO", ""]))
            linha.update(campos)
            aba.append([linha[c] for c in colunas])
        aba.append([])
        aba.append(["Total de operações:", len(linhas)])
        memoria = BytesIO()
        livro.save(memoria)
        livro.close()
        mensagem = EmailMessage()
        mensagem["Subject"] = "Relatório sintético"
        mensagem.set_content("Não seguir links; anexo local.")
        mensagem.add_attachment(memoria.getvalue(), maintype="application",
                               subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=anexo)
        (self.dados / nome).write_bytes(mensagem.as_bytes())

    def rodar(self, competencia=None):
        r = executar(self.dados, self.regras, competencia)
        m = publicar(r, self.saida, "alvo")
        verificar(self.dados, self.saida, self.regras)
        return r, m

    def test_mime_nome_livre_meio_centavo_e_liquidacao_futura(self):
        r, m = self.rodar()
        p = r.partes["REMESSAS"]
        self.assertEqual(p.situacao, "COMPLETO")
        self.assertEqual(p.linhas["comissoes_por_lojista.csv"][0]["comissao_brl"], "131.01")
        self.assertEqual(p.rastro[0]["id_saida"], "REMESSA:2032-04:01234567800012")
        self.assertIn("chegou_agora.eml::nome_livre.xlsx::Operações", p.rastro[0]["arquivo_origem"])
        self.assertTrue(m["partes"]["NOTAS_REMESSA"]["implementado"])
        self.assertEqual(m["resultados"]["notas_a_emitir.csv"]["linhas"], 0)

    def test_tolerancia_abaixo_igual_acima_em_pontos_percentuais(self):
        self.email([{"Spread %": .01509999}, {"Spread %": .0151}, {"Spread %": .01510001, "Valor ME": 100000}])
        r, _ = self.rodar()
        p = r.partes["DIVERGENCIAS_REMESSA"]
        self.assertEqual(len(p.linhas["divergencias_taxa.csv"]), 1)
        self.assertEqual(p.linhas["divergencias_taxa.csv"][0]["ref_origem"], "PC-TESTE-3")
        self.assertEqual(p.linhas["divergencias_taxa.csv"][0]["diferenca_brl"], "50.01")

    def test_percentual_valor_cheio_e_participacao_pelo_realizado(self):
        self.email([{"Spread %": .012, "Spread R$": 17.13}])
        r, _ = self.rodar()
        self.assertEqual(r.partes["REMESSAS"].linhas["comissoes_por_lojista.csv"][0]["comissao_brl"], "8.57")
        self.assertEqual(r.partes["DIVERGENCIAS_REMESSA"].linhas["divergencias_taxa.csv"][0]["diferenca_brl"], "-1.50")

    def test_negativo_irregular_fora_da_comissao_dois_valores(self):
        self.email([{}, {"Spread R$": -10, "Spread %": -.02, "PTAX Ref.": 2, "Valor ME": 100,
                         "Observação": "Taxa travada no pedido (D-1)"}])
        r, _ = self.rodar()
        p = r.partes["REMESSAS"]
        linha = next(l for l in p.linhas["pendencias.csv"] if l["tipo"] == "CONTESTACAO_SPREAD_NEGATIVO")
        self.assertEqual((linha["valor_brl"], linha["valor_reportado_brl"], linha["valor_solicitado_brl"]), ("-13.00", "-10.00", "-13.00"))
        self.assertEqual(p.linhas["comissoes_por_lojista.csv"][0]["comissao_brl"], "131.01")
        self.assertFalse(r.partes["DIVERGENCIAS_REMESSA"].linhas["divergencias_taxa.csv"])

    def test_duplicata_reprocessada_calcula_uma_vez_e_preserva_linha_fisica(self):
        self.email([{"Nº Operação": "PC-TESTE-1"}, {"Nº Operação": "PC-TESTE-1", "Observação": "reprocessada"}])
        r, _ = self.rodar()
        p = r.partes["REMESSAS"]
        self.assertEqual(p.linhas["comissoes_por_lojista.csv"][0]["comissao_brl"], "131.01")
        duplicada = next(l for l in p.linhas["pendencias.csv"] if l["tipo"] == "RELATORIO_PONTE_DUPLICADO")
        self.assertEqual(duplicada["linha_origem"], "4")
        self.assertEqual(p.cobertura, dict(linhas_lidas=2, calculadas=1, pendentes=1, bloqueadas=0))

    def test_duplicata_conflitante_nao_escolhe_dinheiro(self):
        self.email([{"Nº Operação": "PC-TESTE-1"}, {"Nº Operação": "PC-TESTE-1", "Spread R$": 900}])
        r, _ = self.rodar()
        self.assertFalse(r.partes["REMESSAS"].linhas["comissoes_por_lojista.csv"])
        self.assertEqual(r.partes["REMESSAS"].situacao, "PROVISORIO")
        self.assertEqual(len(r.partes["REMESSAS"].linhas["pendencias.csv"]), 2)

    def test_ajuste_antes_do_negativo_rastro_original_e_d10_sem_dobrar_metade(self):
        self.email([{}, {"Spread R$": -235.74, "Spread %": -.10, "Observação": "AJUSTE REF PC-TESTE-1"}])
        r, _ = self.rodar()
        p = r.partes["REMESSAS"]
        self.assertEqual(p.linhas["comissoes_por_lojista.csv"][0]["comissao_brl"], "13.14")
        ajuste = next(l for l in p.rastro if l["id_saida"].startswith("REMESSA:") and l["ref_origem"] == "PC-TESTE-2")
        self.assertEqual(ajuste["valor"], "-117.87")
        self.assertEqual(ajuste["ref_original"], "PC-TESTE-1")
        self.assertTrue(ajuste["arquivo_evento_original"])
        detalhes = [l for l in p.rastro if l["ref_origem"] == "PC-TESTE-2"]
        self.assertTrue(all(l["tipo_lancamento"] == "AJUSTE" and l["ref_original"] == "PC-TESTE-1" for l in detalhes))
        self.assertTrue(any(l["tipo"] == "AJUSTE_PONTE_D10" for l in p.linhas["pendencias.csv"]))
        self.assertFalse(r.partes["DIVERGENCIAS_REMESSA"].linhas["divergencias_taxa.csv"])

    def test_ajuste_de_outro_mes_sem_arquivo_original_nao_inventa_origem(self):
        self.email([{"Spread R$": -1.01, "Observação": "AJUSTE REF PC-ANTERIOR-1"}])
        r, _ = self.rodar()
        t = next(l for l in r.partes["REMESSAS"].rastro if l["id_saida"].startswith("REMESSA:"))
        self.assertEqual((t["valor"], t["ref_original"], t["arquivo_evento_original"]), ("-0.51", "PC-ANTERIOR-1", ""))

    def test_ressalva_aprovada_nao_exclui(self):
        self.email([{"Status Compliance": "APROVADO C/ RESSALVA"}])
        r, _ = self.rodar()
        self.assertEqual(r.partes["REMESSAS"].cobertura["calculadas"], 1)

    def test_novo_cnpj_calcula_participacao_e_pendencia_comercial(self):
        self.email([{"CNPJ Cliente": "12345678000190"}])
        r, _ = self.rodar()
        p = r.partes["REMESSAS"]
        self.assertEqual(p.situacao, "COMPLETO")
        self.assertEqual(p.linhas["comissoes_por_lojista.csv"][0]["comissao_brl"], "131.01")
        self.assertTrue(any(l["quem_decide"] == "Comercial" for l in p.linhas["pendencias.csv"]))
        self.assertEqual(r.partes["DIVERGENCIAS_REMESSA"].situacao, "PROVISORIO")

    def test_sem_planilha_preserva_realizado_sem_inventar_cadastro_ausente(self):
        (self.dados / "taxas.xlsx").unlink()
        r, m = self.rodar()
        self.assertEqual(m["partes"]["CHECKOUT"]["situacao"], "BLOQUEADO")
        self.assertEqual(r.partes["REMESSAS"].situacao, "COMPLETO")
        self.assertEqual(r.partes["DIVERGENCIAS_REMESSA"].situacao, "BLOQUEADO")
        self.assertFalse(any(l["tipo"] == "CADASTRO_REMESSA_PENDENTE" for l in r.partes["REMESSAS"].linhas["pendencias.csv"]))

    def test_sem_orbe_ponte_independente(self):
        (self.dados / "orbe.csv").unlink()
        r, m = self.rodar()
        self.assertEqual(m["situacao"], "BLOQUEADO")
        self.assertEqual(r.partes["REMESSAS"].situacao, "COMPLETO")
        self.assertEqual(r.partes["DIVERGENCIAS_REMESSA"].situacao, "COMPLETO")

    def test_email_ausente_ou_layout_invalido_preserva_minimo(self):
        self.email([{}], layout=False)
        r, _ = self.rodar()
        self.assertEqual(r.situacao, "COMPLETO")
        self.assertEqual(r.partes["REMESSAS"].situacao, "BLOQUEADO")
        (self.dados / "chegou_agora.eml").unlink()
        r, _ = self.rodar()
        self.assertEqual(r.situacao, "COMPLETO")
        self.assertEqual(r.partes["REMESSAS"].situacao, "BLOQUEADO")

    def test_taxa_aplicada_ausente_so_afeta_comparacao(self):
        self.email([{"Spread %": "inválido"}])
        r, _ = self.rodar()
        self.assertEqual(r.partes["REMESSAS"].situacao, "COMPLETO")
        self.assertEqual(r.partes["DIVERGENCIAS_REMESSA"].situacao, "PROVISORIO")

    def test_rastro_comissao_invalido_nao_bloqueia_divergencia_ou_checkout(self):
        self.email([{"Spread %": .02}])
        r = executar(self.dados, self.regras)
        r.partes["REMESSAS"].rastro[0]["valor"] = "999.99"
        m = publicar(r, self.saida, "alvo")
        self.assertEqual(m["partes"]["REMESSAS"]["situacao"], "BLOQUEADO")
        self.assertEqual(m["partes"]["DIVERGENCIAS_REMESSA"]["situacao"], "COMPLETO")
        self.assertEqual(m["partes"]["CHECKOUT"]["situacao"], "COMPLETO")
        verificar(self.dados, self.saida, self.regras)

    def test_parametro_invalido_pendente_nao_contamina_minimo(self):
        self.regras.write_text(self.regras.read_text(encoding="utf-8").replace('tolerancia_pp="0.01"', 'tolerancia_pp="errado"'), encoding="utf-8")
        r, _ = self.rodar()
        self.assertEqual(r.situacao, "COMPLETO")
        self.assertEqual(r.partes["REMESSAS"].situacao, "BLOQUEADO")

    def test_auditor_rejeita_dois_valores_do_negativo_adulterados(self):
        self.email([{"Spread R$": -10}])
        r = executar(self.dados, self.regras)
        p = r.partes["REMESSAS"]
        p.linhas["pendencias.csv"][0]["valor_reportado_brl"] = "-20.00"
        publicar(r, self.saida, "alvo")
        with self.assertRaises(ErroAceite):
            verificar(self.dados, self.saida, self.regras)

    def test_mais_de_um_mes_exige_selecao_e_filtra_pela_contratacao(self):
        self.email([{}, {"Data Contratação": datetime(2032, 5, 1)}])
        r, _ = self.rodar()
        self.assertEqual(r.partes["REMESSAS"].situacao, "BLOQUEADO")
        r, _ = self.rodar("2032-04")
        self.assertEqual(r.partes["REMESSAS"].cobertura["linhas_lidas"], 1)
        self.assertEqual(r.partes["REMESSAS"].situacao, "COMPLETO")

    def test_linha_invalida_nao_vira_zero_ou_operacao_silenciosamente_descartada(self):
        self.email([{}, {"Spread R$": "inexistente"}])
        r, _ = self.rodar()
        p = r.partes["REMESSAS"]
        self.assertEqual(p.situacao, "PROVISORIO")
        self.assertEqual(p.cobertura, dict(linhas_lidas=2, calculadas=1, pendentes=1, bloqueadas=0))
        self.assertTrue(any(l["valor_brl"] == "" for l in p.linhas["pendencias.csv"]))

    def test_duplicata_com_uma_linha_ilegivel_impede_escolha_financeira(self):
        self.email([{"Nº Operação": "PC-TESTE-1"}, {"Nº Operação": "PC-TESTE-1", "Spread R$": "ilegível"}])
        r, _ = self.rodar()
        self.assertEqual(r.partes["REMESSAS"].cobertura["calculadas"], 0)
        self.assertEqual(r.partes["REMESSAS"].cobertura["pendentes"], 2)

    def test_negativo_sem_contrato_pedido_vazio_mas_comissao_disponivel(self):
        self.email([{}, {"Spread R$": -10}])
        (self.dados / "taxas.xlsx").unlink()
        r, _ = self.rodar()
        p = r.partes["REMESSAS"]
        self.assertEqual(p.situacao, "PROVISORIO")
        self.assertEqual(p.linhas["comissoes_por_lojista.csv"][0]["comissao_brl"], "131.01")
        neg = next(l for l in p.linhas["pendencias.csv"] if l["tipo"] == "CONTESTACAO_SPREAD_NEGATIVO")
        self.assertEqual((neg["valor_brl"], neg["valor_reportado_brl"], neg["valor_solicitado_brl"]), ("", "-10.00", ""))

    def test_spread_padrao_em_branco_e_texto_padrao(self):
        for valor in (None, "padrão"):
            with self.subTest(valor=valor):
                livro = load_workbook(self.dados / "taxas.xlsx")
                livro.active.cell(2, 9).value = valor
                livro.save(self.dados / "taxas.xlsx")
                livro.close()
                r, _ = self.rodar()
                self.assertEqual(r.partes["DIVERGENCIAS_REMESSA"].situacao, "COMPLETO")
                self.assertFalse(r.partes["DIVERGENCIAS_REMESSA"].linhas["divergencias_taxa.csv"])

    def test_spread_contratado_zero_nao_vira_padrao(self):
        livro = load_workbook(self.dados / "taxas.xlsx")
        livro.active.cell(2, 9, 0)
        livro.save(self.dados / "taxas.xlsx")
        livro.close()
        self.email([{"Spread %": 0, "Spread R$": 0}])
        r, _ = self.rodar()
        self.assertEqual(r.partes["DIVERGENCIAS_REMESSA"].situacao, "COMPLETO")
        self.assertFalse(r.partes["DIVERGENCIAS_REMESSA"].linhas["divergencias_taxa.csv"])
        self.assertEqual(r.partes["REMESSAS"].linhas["comissoes_por_lojista.csv"][0]["comissao_brl"], "0.00")

    def test_ajuste_malformado_ou_autorreferente_e_pendente(self):
        for obs in ("AJUSTE REF", "AJUSTE REF PC-TESTE-1"):
            self.email([{"Observação": obs, "Spread R$": -10}])
            r, _ = self.rodar()
            self.assertEqual(r.partes["REMESSAS"].situacao, "PROVISORIO")
            self.assertFalse(r.partes["REMESSAS"].linhas["comissoes_por_lojista.csv"])

    def test_planilha_corrompida_nao_impede_realizado(self):
        (self.dados / "taxas.xlsx").write_bytes(b"nao e um XLSX")
        r, _ = self.rodar()
        self.assertEqual(r.partes["REMESSAS"].situacao, "COMPLETO")
        self.assertEqual(r.partes["DIVERGENCIAS_REMESSA"].situacao, "BLOQUEADO")

    def test_auditor_detecta_perda_de_operacao_mesmo_com_somas_reconciliadas(self):
        r = executar(self.dados, self.regras)
        p = r.partes["REMESSAS"]
        p.linhas = {"comissoes_por_lojista.csv": [], "pendencias.csv": []}
        p.rastro = []
        publicar(r, self.saida, "alvo")
        with self.assertRaises(ErroAceite):
            verificar(self.dados, self.saida, self.regras)

    def test_ajuste_abate_mes_do_novo_relatorio_mesmo_com_data_anterior(self):
        self.email([{"Spread R$": -235.74, "Data Contratação": datetime(2032, 3, 2),
                     "Observação": "AJUSTE REF PC-MARCO-1"}], periodo="04/2032")
        r, _ = self.rodar()
        p = r.partes["REMESSAS"]
        self.assertEqual(p.linhas["comissoes_por_lojista.csv"][0]["id_saida"], "REMESSA:2032-04:01234567800012")
        self.assertEqual(p.linhas["comissoes_por_lojista.csv"][0]["comissao_brl"], "-117.87")

