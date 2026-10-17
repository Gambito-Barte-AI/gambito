"""Limites financeiros e identidade da Entrega 2, com dados sintéticos."""

import csv
from datetime import date
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook

from tarvo.execucao import executar
from tarvo.exportacao import publicar
from tarvo.leitores.taxas import vigencia


class Entrega2Test(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.raiz = Path(self.tmp.name)
        self.dados = self.raiz / "dados"
        self.dados.mkdir()
        self.config = self.raiz / "regras.toml"
        self.configurar()
        self.planilha()
        self.vendas()

    def configurar(self, extra=""):
        self.config.write_text('''versao = 1
fonte = "Amostra sintética"
[entrada]
aba_taxas = "Vigente"
abas_identidade = ["Histórica"]
[modalidades]
DEBITO = "debito"
"CREDITO A VISTA" = "credito_vista"
PIX = "pix"
"CREDITO PARCELADO" = "parcelado"
''' + extra, encoding="utf-8")

    def planilha(self, linhas=None, historicas=None):
        livro = Workbook()
        folha = livro.active
        folha.title = "Vigente"
        folha.append(["Cabeçalho deslocado"])
        folha.append(["Lojista", "CNPJ", "EC Orbe", "Débito", "Crédito à vista",
                      "Créd. 2-6x", "Créd. 7-12x", "Pix", "Vigência", "Obs"])
        if linhas is None:
            linhas = [self.contrato()]
        for linha in linhas:
            folha.append(linha)
        antiga = livro.create_sheet("Histórica")
        antiga.append(["Lojista", "CNPJ", "EC Orbe", "Débito", "Crédito à vista"])
        for linha in historicas or []:
            antiga.append(linha)
        livro.save(self.dados / "taxas.xlsx")
        livro.close()

    def contrato(self, ec="42", cnpj="01234567800012", vigencia="01/01/2031", obs="", taxa_vista="2,00%"):
        return ["Loja", cnpj, ec, .01, taxa_vista, .03, "4,00%", .005, vigencia, obs]

    def vendas(self, linhas=None):
        colunas = ["id_transacao", "data_movimento", "cod_estabelecimento", "tipo_lancamento",
                   "modalidade", "qtd_parcelas", "valor_bruto", "id_referencia", "data_venda_original"]
        with (self.dados / "orbe.csv").open("w", encoding="utf-8", newline="") as f:
            escritor = csv.DictWriter(f, fieldnames=colunas, delimiter=";")
            escritor.writeheader()
            for n, linha in enumerate(linhas or [{}], 1):
                escritor.writerow(dict(id_transacao=f"T{n}", data_movimento="15/03/2031",
                                       cod_estabelecimento="42", tipo_lancamento="VENDA",
                                       modalidade="CREDITO A VISTA", qtd_parcelas="1",
                                       valor_bruto="100,00", id_referencia="", data_venda_original="") | linha)

    def executar(self):
        return executar(self.dados, self.config)

    def vinculo(self, ec="42", cnpj="01234567800012", id="VINCULO"):
        return f'''\n[[vinculos]]
id = "{id}"
ec = "{ec}"
cnpj = "{cnpj}"
fonte = "Amostra sintética"
aprovacao = "Teste, sem efeito em produção"
aprovador = "Teste sintético"
aprovado_em = "2031-01-01"
natureza = "INTERPRETACAO_TECNICA"
'''

    def isencao(self, id="ISENCAO"):
        return f'''\n[[isencoes]]
id = "{id}"
cnpj = "01234567800012"
inicio = 2031-03-10
fim = 2031-03-20
fonte = "Amostra sintética"
aprovacao = "Teste, sem efeito em produção"
aprovador = "Teste sintético"
aprovado_em = "2031-01-01"
natureza = "INTERPRETACAO_TECNICA"
'''

    def test_faixas_e_limites_de_parcelas(self):
        self.vendas([{"modalidade": "CREDITO PARCELADO", "qtd_parcelas": str(p)}
                     for p in (2, 6, 7, 12, 13, 1, 0)])
        r = self.executar()
        self.assertEqual([c.valor for c in r.contribuicoes], [Decimal(3), Decimal(3), Decimal(4), Decimal(4)])
        self.assertEqual([c.faixa for c in r.contribuicoes], ["parcelado_2_6"] * 2 + ["parcelado_7_12"] * 2)
        self.assertEqual(len(r.pendencias), 3)
        self.assertTrue(r.validado)

    def test_demais_modalidades(self):
        self.vendas([{"modalidade": m} for m in ("DEBITO", "PIX", "CREDITO A VISTA")])
        self.assertEqual([c.valor for c in self.executar().contribuicoes], [Decimal(1), Decimal("0.50"), Decimal(2)])

    def test_vigencias_por_data_e_limites_inclusivos(self):
        self.planilha([self.contrato(vigencia="até 14/03/2031"),
                       self.contrato(vigencia="a partir de 15/03/2031", taxa_vista="1,50%")])
        self.vendas([{"data_movimento": d} for d in ("14/03/2031", "15/03/2031", "16/03/2031")])
        r = self.executar()
        self.assertEqual([c.valor for c in r.contribuicoes], [Decimal(2), Decimal("1.50"), Decimal("1.50")])
        self.assertEqual([c.contrato.linha for c in r.contribuicoes], [3, 4, 4])

    def test_lacuna_nao_preenche_taxa(self):
        self.planilha([self.contrato(vigencia="até 14/03/2031"),
                       self.contrato(vigencia="a partir de 16/03/2031")])
        self.assertEqual(len(self.executar().pendencias), 1)

    def test_sobreposicao_nao_escolhe_mais_recente(self):
        self.planilha([self.contrato(), self.contrato(vigencia="a partir de 15/03/2031", taxa_vista="99%")])
        r = self.executar()
        self.assertEqual(r.contribuicoes, [])
        self.assertIn("sobrepostas", r.pendencias[0].motivo)

    def test_vigencia_desconhecida_nao_e_descartada(self):
        self.planilha([self.contrato(), self.contrato(vigencia="combinar com Comercial")])
        self.assertEqual(self.executar().contribuicoes, [])

    def test_parser_vigencia_rejeita_data_impossivel(self):
        for texto in ("até 31/02/2031", "a partir de data indefinida", ""):
            self.assertTrue(vigencia(texto)[2])
        self.assertEqual(vigencia(date(2031, 3, 1)), (date(2031, 3, 1), None, ""))

    def test_ec_complementar_busca_por_cnpj_sem_nome(self):
        self.planilha([self.contrato(ec=None)])
        self.configurar(self.vinculo())
        r = self.executar()
        self.assertEqual(r.contribuicoes[0].vinculo, "VINCULO")
        self.assertEqual(r.contribuicoes[0].cnpj, "01234567800012")

    def test_ecs_semelhantes_continuam_separados(self):
        a, b = self.contrato(ec=None), self.contrato(ec=None, cnpj="11234567800012", taxa_vista="5%")
        a[0], b[0] = "Horizonte Um", "Horizonte Dois"
        self.planilha([a, b])
        self.configurar(self.vinculo() + self.vinculo("43", "11234567800012", "OUTRO"))
        self.vendas([{}, {"cod_estabelecimento": "43"}])
        r = self.executar()
        self.assertEqual([c.cnpj for c in r.contribuicoes], ["01234567800012", "11234567800012"])
        self.assertEqual([c.valor for c in r.contribuicoes], [Decimal(2), Decimal(5)])

    def test_conflito_entre_vinculo_e_planilha_nao_sobrescreve(self):
        self.configurar(self.vinculo(cnpj="11234567800012"))
        r = self.executar()
        self.assertEqual(r.contribuicoes, [])
        self.assertIn("conflitante", r.pendencias[0].motivo)

    def test_conflito_entre_vinculos(self):
        self.configurar(self.vinculo() + self.vinculo(cnpj="11234567800012", id="OUTRO"))
        self.assertEqual(self.executar().contribuicoes, [])

    def test_cnpj_invalido_na_planilha_nao_e_substituido_por_configuracao(self):
        self.planilha([self.contrato(cnpj="inválido"), self.contrato(ec=None)])
        self.configurar(self.vinculo())
        self.assertEqual(self.executar().contribuicoes, [])

    def test_identidade_historica_nunca_fornece_taxa(self):
        self.planilha([], [["Histórica", "01234567800012", "42", .99, .99]])
        self.configurar(self.vinculo())
        r = self.executar()
        self.assertEqual(r.contribuicoes, [])
        self.assertEqual(r.pendencias[0].cnpj, "01234567800012")
        self.assertEqual(r.pendencias[0].identidade_origem.ref, "Histórica!L2")
        manifesto = publicar(r, self.raiz / "saida")
        self.assertEqual(manifesto["pendencias"][0]["ref_identidade"], "Histórica!L2")

    def test_taxa_antiga_nao_substitui_taxa_vigente_ausente(self):
        self.planilha([self.contrato(taxa_vista=None)], [["Antiga", "01234567800012", "42", .99, .99]])
        self.assertEqual(self.executar().contribuicoes, [])

    def test_cnpj_numerico_e_mascarado_normalizam_para_mesma_chave(self):
        for cnpj in (1234567800012, "01.234.567/8000-12"):
            self.planilha([self.contrato(cnpj=cnpj)])
            self.assertEqual(self.executar().contribuicoes[0].cnpj, "01234567800012")

    def test_identidade_historica_nao_exige_aba_opcional(self):
        livro = Workbook()
        folha = livro.active
        folha.title = "Vigente"
        folha.append(["Lojista", "CNPJ", "EC Orbe", "Débito", "Crédito à vista",
                      "Créd. 2-6x", "Créd. 7-12x", "Pix", "Vigência", "Obs"])
        folha.append(self.contrato())
        livro.save(self.dados / "taxas.xlsx")
        livro.close()
        self.assertEqual(self.executar().situacao, "COMPLETO")

    def test_sem_coluna_parcelada_nao_reaproveita_taxa_vista(self):
        self.planilha([self.contrato()])
        from openpyxl import load_workbook
        livro = load_workbook(self.dados / "taxas.xlsx")
        livro["Vigente"].delete_cols(6, 2)
        livro.save(self.dados / "taxas.xlsx")
        livro.close()
        self.vendas([{"modalidade": "CREDITO PARCELADO", "qtd_parcelas": "6"}])
        r = self.executar()
        self.assertEqual(r.contribuicoes, [])
        self.assertIn("Taxa contratada ausente", r.pendencias[0].motivo)

    def test_bloqueio_comercial_prevalece_sobre_planilha(self):
        self.configurar('''\n[[contratos_bloqueados]]
id = "BLOQUEIO"
cnpj = "01234567800012"
motivo = "Aguardando confirmação comercial; decisão de apresentação pendente"
fonte = "Amostra"
aprovacao = "Teste"
aprovador = "Teste sintético"
aprovado_em = "2031-01-01"
natureza = "INTERPRETACAO_TECNICA"
''')
        r = self.executar()
        self.assertEqual(r.contribuicoes, [])
        self.assertEqual(r.pendencias[0].quem_decide, "Comercial")
        self.vendas([{"tipo_lancamento": "ANTECIPACAO"}])
        self.assertEqual(self.executar().pendencias[0].quem_decide, "Comercial")

    def test_isencao_por_periodo_e_nao_por_mes_inteiro(self):
        self.configurar(self.isencao())
        self.vendas([{"data_movimento": d} for d in ("09/03/2031", "10/03/2031", "20/03/2031", "21/03/2031")])
        r = self.executar()
        self.assertEqual([c.valor for c in r.contribuicoes], [Decimal(2), Decimal(0), Decimal(0), Decimal(2)])
        self.assertEqual([c.regra for c in r.contribuicoes], ["PLANILHA", "ISENCAO", "ISENCAO", "PLANILHA"])
        self.assertEqual(len(r.rastro), 4)  # contribuições zero também têm origem

    def test_isencao_nao_esconde_sobreposicao_contratual(self):
        self.configurar(self.isencao())
        self.planilha([self.contrato(), self.contrato()])
        self.assertEqual(self.executar().contribuicoes, [])

    def test_isencoes_sobrepostas_viram_pendencia(self):
        self.configurar(self.isencao() + self.isencao("OUTRA"))
        self.assertIn("Isenções sobrepostas", self.executar().pendencias[0].motivo)

    def test_isencao_nao_ignora_nova_observacao(self):
        self.configurar(self.isencao())
        self.planilha([self.contrato(obs="nova condição não confirmada")])
        self.assertEqual(self.executar().contribuicoes, [])

    def test_observacao_esclarecida_e_restrita_ao_cnpj_e_texto(self):
        self.configurar('''\n[[observacoes_esclarecidas]]
id = "OBS"
cnpj = "01234567800012"
texto = "EC já confirmado"
fonte = "Amostra"
aprovacao = "Teste"
aprovador = "Teste sintético"
aprovado_em = "2031-01-01"
natureza = "INTERPRETACAO_TECNICA"
''')
        self.planilha([self.contrato(obs="EC já confirmado")])
        self.assertEqual(len(self.executar().contribuicoes), 1)
        self.planilha([self.contrato(obs="EC já confirmado; tarifa especial")])
        self.assertEqual(self.executar().contribuicoes, [])

    def test_regras_invalidas_bloqueiam_antes_do_calculo(self):
        for extra in (self.vinculo().replace('"01234567800012"', '"123"'),
                      self.isencao().replace("fim = 2031-03-20", "fim = 2031-03-01"),
                      self.vinculo().replace('fonte = "Amostra sintética"', 'fonte = ""')):
            self.configurar(extra)
            self.assertEqual(self.executar().situacao, "BLOQUEADO")

    def test_eventos_sem_sinal_referencia_ou_taxa_continuam_pendentes(self):
        self.vendas([{"tipo_lancamento": tipo} for tipo in ("ESTORNO", "ANTECIPACAO", "REVERSAO ANTECIPACAO")])
        r = self.executar()
        self.assertEqual(r.contribuicoes, [])
        self.assertEqual(len(r.pendencias), 3)


if __name__ == "__main__":
    unittest.main()
