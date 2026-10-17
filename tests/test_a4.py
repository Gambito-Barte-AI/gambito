"""A4: propostas por saldo, rastros assinados, isolamento e regressão agnóstica."""

import csv
from decimal import Decimal
import unittest

from openpyxl import load_workbook

import test_a3 as amostras_ponte
from validar_saida import ErroAceite, verificar
from tarvo.execucao import executar
from tarvo.exportacao import publicar
from tarvo.modelos import ResultadoParte
from tarvo.notas import calcular
from tarvo.validacao import validar_parte


class A4Test(unittest.TestCase):
    vendas = amostras_ponte.A3Test.vendas
    email = amostras_ponte.A3Test.email

    def setUp(self):
        amostras_ponte.A3Test.setUp(self)
        with self.regras.open("a", encoding="utf-8") as f:
            for tabela in ("notas_checkout", "notas_remessa"):
                f.write(f'''\n[{tabela}]
id="{tabela}_TESTE"
servico="Serviço sintético {tabela}"
fonte="Teste sintético, sem vínculo aos exemplos oficiais"
aprovacao="Amostra para teste"
aprovador="Teste"
aprovado_em="NAO_INFORMADO"
natureza="INTERPRETACAO_TECNICA"
''')
                if tabela == "notas_remessa":
                    f.write('tomador_cnpj="00987654000132"\ntomador="Contraparte sintética"\n')

    def rodar(self, competencia=None):
        r = executar(self.dados, self.regras, competencia)
        m = publicar(r, self.saida, "alvo")
        verificar(self.dados, self.saida, self.regras)
        return r, m

    def orbe(self, linhas):
        colunas = ["id_transacao", "data_movimento", "cod_estabelecimento", "tipo_lancamento",
                   "modalidade", "qtd_parcelas", "valor_bruto", "id_referencia", "data_venda_original"]
        with (self.dados / "orbe.csv").open("w", encoding="utf-8", newline="") as f:
            escritor = csv.DictWriter(f, fieldnames=colunas, delimiter=";")
            escritor.writeheader()
            for campos in linhas:
                linha = dict.fromkeys(colunas, "")
                linha.update(id_transacao="VENDA", data_movimento="02/04/2032", cod_estabelecimento="42",
                             tipo_lancamento="VENDA", modalidade="PIX", qtd_parcelas="1", valor_bruto="100,00")
                linha.update(campos)
                escritor.writerow(linha)

    def test_checkout_positivo_e_nota_unica_da_contraparte_configurada(self):
        r, m = self.rodar()
        a = r.partes["NOTAS_CHECKOUT"].linhas["notas_a_emitir.csv"]
        b = r.partes["NOTAS_REMESSA"].linhas["notas_a_emitir.csv"]
        self.assertEqual((len(a), a[0]["valor_brl"], a[0]["tomador_cnpj"]), (1, "0.01", "01234567800012"))
        self.assertEqual((len(b), b[0]["valor_brl"], b[0]["tomador_cnpj"]), (1, "131.01", "00987654000132"))
        self.assertEqual(b[0]["tomador"], "Contraparte sintética")
        self.assertIn("2032-04", b[0]["discriminacao"])
        self.assertEqual(m["resultados"]["notas_a_emitir.csv"]["situacao"], "COMPLETO")
        self.assertTrue(m["partes"]["EXTRATO_AURORA"]["implementado"])
        self.assertEqual(m["partes"]["EXTRATO_AURORA"]["situacao"], "BLOQUEADO")

    def test_checkout_nao_agrupa_por_nome_do_estabelecimento(self):
        livro = load_workbook(self.dados / "taxas.xlsx")
        nome = livro.active.cell(2, 1).value
        livro.active.append([nome, "12345678000190", "84", .01, .02, .005, "01/01/2030", "", .015])
        livro.save(self.dados / "taxas.xlsx")
        livro.close()
        self.orbe([{}, dict(id_transacao="OUTRA", cod_estabelecimento="84")])
        r, _ = self.rodar()
        notas = r.partes["NOTAS_CHECKOUT"].linhas["notas_a_emitir.csv"]
        self.assertEqual({n["tomador_cnpj"] for n in notas}, {"01234567800012", "12345678000190"})

    def test_checkout_zero_sem_nota_com_pendencia_e_rastro_zero_explicito(self):
        self.orbe([dict(valor_bruto="0,01")])
        r, _ = self.rodar()
        p = r.partes["NOTAS_CHECKOUT"]
        self.assertFalse(p.linhas["notas_a_emitir.csv"])
        self.assertEqual(p.linhas["pendencias.csv"][0]["valor_brl"], "0.00")
        self.assertEqual(p.rastro[0]["valor"], "0.00")
        self.assertEqual(p.situacao, "COMPLETO")

    def test_checkout_negativo_sem_nota_ou_transporte_ao_mes_seguinte(self):
        self.orbe([dict(id_transacao="E", tipo_lancamento="ESTORNO", valor_bruto="-100,00",
                        id_referencia="HISTORICA", data_venda_original="02/03/2032")])
        r, _ = self.rodar()
        p = r.partes["NOTAS_CHECKOUT"]
        self.assertFalse(p.linhas["notas_a_emitir.csv"])
        self.assertEqual(p.linhas["pendencias.csv"][0]["valor_brl"], "-0.50")
        self.vendas(data="02/05/2032")
        (self.dados / "chegou_agora.eml").unlink()
        r, _ = self.rodar()
        self.assertEqual(r.partes["NOTAS_CHECKOUT"].linhas["notas_a_emitir.csv"][0]["valor_brl"], "0.01")
        self.assertFalse(r.partes["NOTAS_CHECKOUT"].linhas["pendencias.csv"])

    def test_ajuste_da_ponte_abate_nota_sem_aplicar_metade_novamente(self):
        self.email([{}, {"Spread R$": -235.74, "Observação": "AJUSTE REF PC-TESTE-1"}])
        r, _ = self.rodar()
        p = r.partes["NOTAS_REMESSA"]
        self.assertEqual(p.linhas["notas_a_emitir.csv"][0]["valor_brl"], "13.14")
        self.assertEqual([l["valor"] for l in p.rastro], ["131.01", "-117.87"])
        self.assertEqual(p.rastro[1]["ref_original"], "PC-TESTE-1")
        self.assertEqual(p.rastro[1]["arquivo_evento_original"], p.rastro[0]["arquivo_origem"])
        self.assertIn("D-10", p.linhas["notas_a_emitir.csv"][0]["discriminacao"])
        self.assertTrue(any("D-10" in m for m in p.motivos))

    def test_ponte_saldo_zero_ou_negativo_vira_pendencia_unica(self):
        for ajuste, saldo in ((-2, "0.00"), (-4, "-1.00")):
            with self.subTest(ajuste=ajuste):
                self.email([{"Spread R$": 2}, {"Spread R$": ajuste, "Observação": "AJUSTE REF PC-TESTE-1"}])
                r, _ = self.rodar()
                p = r.partes["NOTAS_REMESSA"]
                self.assertFalse(p.linhas["notas_a_emitir.csv"])
                self.assertEqual(len(p.linhas["pendencias.csv"]), 1)
                self.assertEqual(p.linhas["pendencias.csv"][0]["valor_brl"], saldo)

    def test_ponte_total_unico_inclui_lojista_com_contribuicao_negativa(self):
        self.email([{"Spread R$": 10}, {"CNPJ Cliente": "12345678000190", "Spread R$": -4,
                                       "Observação": "AJUSTE REF PC-OUTRO-MES"}])
        r, _ = self.rodar()
        p = r.partes["NOTAS_REMESSA"]
        self.assertEqual(len(p.linhas["notas_a_emitir.csv"]), 1)
        self.assertEqual(p.linhas["notas_a_emitir.csv"][0]["valor_brl"], "3.00")
        self.assertFalse(p.linhas["pendencias.csv"])
        self.assertEqual(len(p.rastro), 2)

    def test_pendencias_divergencias_duplicatas_e_negativo_irregular_nao_somam_na_nota(self):
        self.email([{"Nº Operação": "PC-TESTE-1", "Spread %": .02},
                    {"Nº Operação": "PC-TESTE-1", "Spread %": .02, "Observação": "reprocessada"},
                    {"Spread R$": -10}])
        r, _ = self.rodar()
        p = r.partes["NOTAS_REMESSA"]
        self.assertEqual(p.linhas["notas_a_emitir.csv"][0]["valor_brl"], "131.01")
        self.assertEqual(len(p.rastro), 1)
        self.assertEqual(p.rastro[0]["ref_origem"], "PC-TESTE-1")

    def test_sem_ponte_notas_checkout_continuam(self):
        (self.dados / "chegou_agora.eml").unlink()
        r, _ = self.rodar()
        self.assertEqual(r.partes["NOTAS_CHECKOUT"].situacao, "COMPLETO")
        self.assertEqual(r.partes["NOTAS_REMESSA"].situacao, "BLOQUEADO")

    def test_sem_orbe_notas_remessa_continuam(self):
        (self.dados / "orbe.csv").unlink()
        r, _ = self.rodar()
        self.assertEqual(r.partes["NOTAS_CHECKOUT"].situacao, "BLOQUEADO")
        self.assertEqual(r.partes["NOTAS_REMESSA"].situacao, "COMPLETO")

    def test_divergencia_corrompida_nao_bloqueia_nota_da_comissao(self):
        self.email([{"Spread %": .02}])
        r = executar(self.dados, self.regras)
        r.partes["DIVERGENCIAS_REMESSA"].rastro[0]["valor"] = "999.99"
        m = publicar(r, self.saida, "alvo")
        self.assertEqual(m["partes"]["DIVERGENCIAS_REMESSA"]["situacao"], "BLOQUEADO")
        self.assertEqual(m["partes"]["NOTAS_REMESSA"]["situacao"], "COMPLETO")
        verificar(self.dados, self.saida, self.regras)

    def test_rastro_comissao_invalido_bloqueia_somente_suas_notas(self):
        r = executar(self.dados, self.regras)
        r.rastro[0]["valor"] = "9.99"
        m = publicar(r, self.saida, "alvo")
        self.assertEqual(m["partes"]["NOTAS_CHECKOUT"]["situacao"], "BLOQUEADO")
        self.assertEqual(m["partes"]["NOTAS_REMESSA"]["situacao"], "COMPLETO")
        self.assertEqual(m["partes"]["REMESSAS"]["situacao"], "COMPLETO")
        verificar(self.dados, self.saida, self.regras)

    def test_nota_adulterada_com_soma_reconciliada_bloqueia_so_nota(self):
        r = executar(self.dados, self.regras)
        p = r.partes["NOTAS_CHECKOUT"]
        p.linhas["notas_a_emitir.csv"][0]["valor_brl"] = "9.99"
        p.rastro[0]["valor"] = "9.99"
        m = publicar(r, self.saida, "alvo")
        self.assertEqual(m["partes"]["NOTAS_CHECKOUT"]["situacao"], "BLOQUEADO")
        self.assertEqual(m["partes"]["CHECKOUT"]["situacao"], "COMPLETO")
        self.assertEqual(m["partes"]["NOTAS_REMESSA"]["situacao"], "COMPLETO")
        verificar(self.dados, self.saida, self.regras)

    def test_tomador_remessa_invalido_nao_bloqueia_outros_resultados(self):
        self.regras.write_text(self.regras.read_text(encoding="utf-8").replace('tomador_cnpj="00987654000132"', 'tomador_cnpj="inválido"'), encoding="utf-8")
        r, _ = self.rodar()
        self.assertEqual(r.partes["NOTAS_REMESSA"].situacao, "BLOQUEADO")
        self.assertEqual(r.partes["NOTAS_CHECKOUT"].situacao, "COMPLETO")
        self.assertEqual(r.partes["REMESSAS"].situacao, "COMPLETO")

    def test_proposta_herda_situacao_provisoria_do_subconjunto_calculado(self):
        self.orbe([{}, dict(id_transacao="INVALIDA", modalidade="NAO_RECONHECIDA")])
        r, _ = self.rodar()
        self.assertEqual(r.partes["NOTAS_CHECKOUT"].situacao, "PROVISORIO")
        self.assertEqual(r.partes["NOTAS_CHECKOUT"].linhas["notas_a_emitir.csv"][0]["valor_brl"], "0.50")

    def test_auditor_rejeita_nota_adulterada_apos_publicacao_mesmo_com_hash_novo(self):
        import json
        from hashlib import sha256
        self.rodar()
        caminho = self.saida / "notas_a_emitir.csv"
        with caminho.open(encoding="utf-8", newline="") as f:
            leitor = csv.DictReader(f, delimiter=";")
            colunas, linhas = leitor.fieldnames, list(leitor)
        linhas[0]["tomador_cnpj"] = "11111111000111"
        with caminho.open("w", encoding="utf-8", newline="") as f:
            escritor = csv.DictWriter(f, fieldnames=colunas, delimiter=";")
            escritor.writeheader(); escritor.writerows(linhas)
        manifesto = self.saida / "manifesto.json"
        m = json.loads(manifesto.read_text(encoding="utf-8"))
        m["arquivos"][caminho.name] = sha256(caminho.read_bytes()).hexdigest()
        manifesto.write_text(json.dumps(m), encoding="utf-8")
        with self.assertRaises(ErroAceite):
            verificar(self.dados, self.saida, self.regras)


class NotasCompetenciaTest(unittest.TestCase):
    def test_datasets_de_outros_meses_e_saldos_nao_se_compensam_entre_competencias(self):
        p = ResultadoParte(implementado=True, situacao="COMPLETO", validado=True,
            relatorios=("comissoes_por_lojista.csv",), linhas={"comissoes_por_lojista.csv": []})
        for mes, valor in (("2040-12", "-2.00"), ("2041-01", "3.00")):
            identificador = f"CHECKOUT:{mes}:11111111000111"
            p.linhas["comissoes_por_lojista.csv"].append(dict(id_saida=identificador, cnpj="11111111000111",
                lojista="Nome sintético livre", linha_negocio="CHECKOUT", comissao_brl=valor))
            p.rastro.append(dict(id_saida=identificador, arquivo_origem="arquivo_livre.csv", ref_origem=mes, valor=valor))
        validar_parte(p)
        n = calcular("NOTAS_CHECKOUT", p, dict(servico="Serviço sintético"))
        validar_parte(n)
        self.assertEqual(n.linhas["notas_a_emitir.csv"][0]["valor_brl"], "3.00")
        self.assertEqual(n.linhas["pendencias.csv"][0]["valor_brl"], "-2.00")
        self.assertIn("2041-01", n.linhas["notas_a_emitir.csv"][0]["discriminacao"])
