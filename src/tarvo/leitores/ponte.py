"""MIME e XLSX em memória; nomes de arquivos mensais não definem o conteúdo."""

from datetime import date, datetime
from dataclasses import replace
from decimal import Decimal, InvalidOperation
from email import policy
from email.parser import BytesParser
from hashlib import sha256
from io import BytesIO
from pathlib import PureWindowsPath
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile
import re

from openpyxl import load_workbook

from .taxas import identificador, normalizar
from ..modelos import ErroEntrada, OperacaoPonte, Origem, Pendencia

COLUNAS = {"no operacao", "data contratacao", "cnpj cliente", "cliente", "spread r$", "observacao", "status compliance"}


def numero(valor, dinheiro=False):
    if isinstance(valor, bool) or valor is None:
        raise ValueError("número ausente ou inválido")
    texto = str(valor).strip()
    percentual = texto.endswith("%")
    if percentual:
        texto = texto[:-1]
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        n = Decimal(texto)
    except InvalidOperation as exc:
        raise ValueError("número inválido") from exc
    if not n.is_finite() or (dinheiro and n != n.quantize(Decimal(".01"))):
        raise ValueError("valor não finito ou fração de centavo monetário")
    return n / 100 if percentual else n


def data(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return datetime.strptime(str(valor), "%d/%m/%Y").date()


def opcional(campos, nome):
    try:
        return numero(campos.get(nome))
    except ValueError:
        return None


def ler_anexo(raw, arquivo):
    operacoes, pendencias, quantidade, reconhecido = [], [], 0, False
    livro = load_workbook(BytesIO(raw), read_only=True, data_only=False)
    try:
        for aba in livro:
            indices = None
            competencia_relatorio = ""
            for linha, celulas in enumerate(aba.iter_rows(), 1):
                valores = [c.value for c in celulas]
                if indices is None:
                    for v in valores:
                        texto = normalizar(v)
                        if texto.startswith("competencia:"):
                            periodo = re.match(r"competencia:\s*(0[1-9]|1[0-2])/(\d{4})\b", texto)
                            if not periodo:
                                raise ErroEntrada(f"{arquivo}: competência declarada inválida")
                            competencia_relatorio = f"{periodo[2]}-{periodo[1]}"
                    nomes = [normalizar(v).replace("º", "o") for v in valores]
                    if COLUNAS <= set(nomes):
                        if any(nomes.count(n) != 1 for n in nomes if n):
                            raise ErroEntrada(f"{arquivo}: cabeçalho duplicado")
                        indices = {n: p for p, n in enumerate(nomes) if n}
                        reconhecido = True
                    continue
                if all(v is None for v in valores):
                    continue
                campos = {n: valores[p] for n, p in indices.items()}
                # Rodapés conhecidos não são operações; qualquer outra linha tem destino.
                if normalizar(campos.get("no operacao")) in {"total de operacoes:", "spread total r$:"}:
                    continue
                quantidade += 1
                ref = str(campos.get("no operacao") or "").strip()
                origem = Origem(f"{arquivo}::{aba.title}", ref or f"L{linha}", linha)
                try:
                    if not ref:
                        raise ValueError("Nº Operação ausente; referência física usada somente no diagnóstico")
                    if any(c.data_type == "f" for c in celulas):
                        raise ValueError("fórmula em operação não suportada")
                    cnpj = identificador(campos["cnpj cliente"], cnpj=True)
                    if not cnpj or not str(campos["cliente"] or "").strip():
                        raise ValueError("CNPJ ou cliente ausente/inválido")
                    spread = numero(campos["spread r$"], dinheiro=True)
                    observacao = str(campos["observacao"] or "").strip()
                    operacoes.append(OperacaoPonte(origem, data(campos["data contratacao"]), cnpj,
                        str(campos["cliente"]).strip(), spread, opcional(campos, "spread %"),
                        opcional(campos, "ptax ref."), opcional(campos, "valor me"),
                        observacao, str(campos["status compliance"] or "").strip(),
                        competencia_relatorio if observacao.upper().startswith("AJUSTE REF") else ""))
                except (ValueError, InvalidOperation) as exc:
                    pendencias.append(Pendencia(origem, f"Operação Ponte inválida: {exc}", "Financeiro"))
    finally:
        livro.close()
    if not reconhecido:
        raise ErroEntrada(f"{arquivo}: cabeçalho de operações Ponte não reconhecido")
    return operacoes, pendencias, quantidade


def ler_pasta(pasta, competencia=None):
    eventos, pendencias, lidas, fontes, avisos = [], [], 0, [], []
    caminhos = sorted(pasta.glob("*.eml"))
    for caminho in caminhos:
        raw = caminho.read_bytes()
        fonte = dict(arquivo=caminho.name, sha256=sha256(raw).hexdigest(), anexos=[])
        fontes.append(fonte)
        try:
            locais, falhas_locais, lidas_locais = [], [], 0
            mensagem = BytesParser(policy=policy.default).parsebytes(raw)
            nomes = set()
            for parte in mensagem.walk():
                nome = PureWindowsPath(parte.get_filename() or "").name
                if not nome.lower().endswith(".xlsx"):
                    continue
                if nome in nomes:
                    raise ErroEntrada("Anexos XLSX com nomes repetidos no mesmo e-mail")
                nomes.add(nome)
                anexo = parte.get_payload(decode=True)
                if not anexo:
                    raise ErroEntrada("Anexo XLSX vazio")
                ops, ps, qtd = ler_anexo(anexo, f"{caminho.name}::{nome}")
                locais.extend(ops); falhas_locais.extend(ps); lidas_locais += qtd
                fonte["anexos"].append(dict(arquivo=nome, sha256=sha256(anexo).hexdigest(), linhas=qtd))
            if not nomes:
                raise ErroEntrada("E-mail sem anexo XLSX")
            eventos.extend(locais); pendencias.extend(falhas_locais); lidas += lidas_locais
            fonte["situacao"] = "DISPONIVEL"
        except (ErroEntrada, ValueError, OSError, KeyError, TypeError, BadZipFile, ParseError) as exc:
            fonte.update(situacao="BLOQUEADO", motivo=str(exc))
            avisos.append(f"{caminho.name}: {exc}")
    if competencia:
        eventos = [replace(e, competencia_ajuste=competencia) if e.observacao.upper().startswith("AJUSTE REF")
                   and not e.competencia_ajuste else e for e in eventos]
    meses = sorted({e.competencia for e in eventos})
    if competencia:
        eventos = [e for e in eventos if e.competencia == competencia]
        # Datas inválidas ficam pendentes, sem atribuir competência por estimativa.
        lidas = len(eventos) + len(pendencias)
    ambiguo = not competencia and len(meses) > 1
    if ambiguo:
        avisos.append("Ponte: múltiplas competências; informe --competencia")
    disponivel = any(f["situacao"] == "DISPONIVEL" for f in fontes)
    estado = "BLOQUEADO" if ambiguo else "PROVISORIO" if disponivel and (avisos or pendencias) else "DISPONIVEL" if disponivel else "BLOQUEADO"
    return eventos, pendencias, lidas, dict(situacao=estado, arquivos=fontes, avisos=avisos,
        competencias=meses, competencia=competencia or (meses[0] if len(meses) == 1 else None),
        motivo="; ".join(avisos) if avisos else "Nenhum e-mail Ponte disponível" if not disponivel else "")
