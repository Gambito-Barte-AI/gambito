"""Contratos vigentes e identidade histórica; taxas antigas nunca são carregadas."""

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook

from ..modelos import Cadastro, Contrato, ErroEntrada, Identidade, Origem

COLUNAS = {"lojista", "cnpj", "ec orbe", "debito", "credito a vista", "pix", "vigencia", "obs"}


def normalizar(valor) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    return " ".join("".join(c for c in texto if not unicodedata.combining(c)).lower().split())


def identificador(valor, cnpj=False) -> str:
    if valor is None or isinstance(valor, bool):
        return ""
    if isinstance(valor, (int, float)):
        decimal = Decimal(str(valor))
        if not decimal.is_finite() or decimal != decimal.to_integral_value():
            return ""
        texto = str(int(decimal))
    else:
        texto = str(valor).strip()
        if cnpj:
            texto = texto.replace(".", "").replace("/", "").replace("-", "")
    if not re.fullmatch(r"[0-9]+", texto):
        return ""
    if cnpj:
        return texto.zfill(14) if len(texto) <= 14 else ""
    return texto


def taxa(valor) -> Decimal | None:
    if valor is None or isinstance(valor, bool):
        return None
    try:
        if isinstance(valor, (int, float)):
            numero = Decimal(str(valor))
        else:
            texto = str(valor).strip()
            if not texto.endswith("%"):
                return None  # texto sem unidade não é interpretado por aproximação
            numero = Decimal(texto[:-1].replace(",", ".")) / Decimal(100)
        return numero if numero.is_finite() and Decimal(0) <= numero <= Decimal(1) else None
    except InvalidOperation:
        return None


def inicio_simples(valor) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    try:
        return datetime.strptime(str(valor).strip(), "%d/%m/%Y").date()
    except ValueError:
        return None


def vigencia(valor) -> tuple[date | None, date | None, str]:
    simples = inicio_simples(valor)
    if simples is not None:
        return simples, None, ""
    texto = normalizar(valor)
    for prefixo, inicio in (("ate ", False), ("a partir de ", True)):
        if texto.startswith(prefixo):
            limite = inicio_simples(texto[len(prefixo):])
            if limite:
                return (limite, None, "") if inicio else (None, limite, "")
    return None, None, "Vigência ausente ou não reconhecida; requer decisão do Comercial"


def descobrir(pasta: Path) -> Path:
    arquivos = sorted(p for p in pasta.glob("*.xlsx") if not p.name.startswith("~$"))
    if len(arquivos) != 1:
        raise ErroEntrada(f"Taxas: esperado um XLSX na pasta; encontrados {len(arquivos)}")
    return arquivos[0]


def ler(caminho: Path, aba: str, abas_identidade: list[str] | None = None) -> Cadastro:
    try:
        livro = load_workbook(caminho, read_only=True, data_only=False)
        try:
            if aba not in livro.sheetnames or "antigo" in normalizar(aba):
                raise ErroEntrada(f"{caminho.name}: aba de taxas vigente inválida/ausente: {aba}")
            linhas = livro[aba].iter_rows()
            indices = None
            cadastro = Cadastro()
            for numero, celulas in enumerate(linhas, 1):
                if indices is None:
                    nomes = [normalizar(c.value) for c in celulas]
                    if COLUNAS <= set(nomes):
                        if any(nomes.count(nome) != 1 for nome in COLUNAS):
                            raise ErroEntrada(f"{caminho.name}: L{numero}: colunas duplicadas")
                        indices = {nome: nomes.index(nome) for nome in COLUNAS}
                        for coluna in ("cred. 2-6x", "cred. 7-12x", "antecipacao"):
                            if nomes.count(coluna) > 1:
                                raise ErroEntrada(f"{caminho.name}: L{numero}: faixa duplicada: {coluna}")
                            if coluna in nomes:
                                indices[coluna] = nomes.index(coluna)
                    continue
                campos = {nome: celulas[pos] for nome, pos in indices.items()}
                if all(c.value is None for c in campos.values()):
                    continue
                ec = identificador(campos["ec orbe"].value)
                # Rodapés sem CNPJ/EC não são linhas contratuais.
                if not ec and campos["cnpj"].value is None:
                    continue
                cnpj = identificador(campos["cnpj"].value, cnpj=True)
                nome = str(campos["lojista"].value or "").strip()
                inicio, fim, problema_vigencia = vigencia(campos["vigencia"].value)
                problemas = []
                if any(c.data_type == "f" for c in campos.values()):
                    problemas.append("Fórmula na linha contratual não suportada")
                if not cnpj or not nome:
                    problemas.append("Identidade contratual inválida")
                if problema_vigencia:
                    problemas.append(problema_vigencia)
                origem = Origem(caminho.name, f"{aba}!L{numero}", numero)
                if cnpj and nome and all(campos[k].data_type != "f" for k in ("cnpj", "lojista", "ec orbe")):
                    cadastro.identidades.append(Identidade(cnpj, nome, origem, ec))
                cadastro.contratos.append(Contrato(
                    origem, ec, cnpj, nome,
                    {chave: taxa(campos[coluna].value) if coluna in campos else None
                     for chave, coluna in (("debito", "debito"), ("credito_vista", "credito a vista"),
                                           ("pix", "pix"), ("parcelado_2_6", "cred. 2-6x"),
                                           ("parcelado_7_12", "cred. 7-12x"),
                                           ("antecipacao", "antecipacao"))},
                    inicio, "; ".join(problemas), fim, str(campos["obs"].value or "").strip(),
                ))
            if indices is None:
                raise ErroEntrada(f"{caminho.name}: cabeçalho não reconhecido na aba {aba}")
            for nome_aba in abas_identidade or []:
                if nome_aba not in livro.sheetnames:
                    continue  # fonte histórica é opcional, nunca bloqueia contratos vigentes
                indices_identidade = None
                for numero, celulas in enumerate(livro[nome_aba].iter_rows(), 1):
                    if indices_identidade is None:
                        nomes = [normalizar(c.value) for c in celulas]
                        if nomes.count("cnpj") == 1 and nomes.count("lojista") == 1:
                            indices_identidade = {k: nomes.index(k) for k in ("cnpj", "lojista")}
                        continue
                    cnpj_celula = celulas[indices_identidade["cnpj"]]
                    nome_celula = celulas[indices_identidade["lojista"]]
                    if cnpj_celula.data_type == "f" or nome_celula.data_type == "f":
                        continue
                    cnpj = identificador(cnpj_celula.value, cnpj=True)
                    nome = str(nome_celula.value or "").strip()
                    if cnpj and nome:
                        cadastro.identidades.append(Identidade(
                            cnpj, nome, Origem(caminho.name, f"{nome_aba}!L{numero}", numero), historica=True,
                        ))
            return cadastro
        finally:
            livro.close()
    except ErroEntrada:
        raise
    except Exception as exc:
        raise ErroEntrada(f"{caminho.name}: XLSX inválido: {exc}") from exc
