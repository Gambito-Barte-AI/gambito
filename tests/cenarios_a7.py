"""Entradas inteiramente sintéticas de outubro/2032 para aceite integrado."""

import csv
from datetime import datetime
from email.message import EmailMessage
from io import BytesIO
import json
from pathlib import Path

from openpyxl import Workbook


def criar(raiz: Path):
    dados = raiz / "dados"
    dados.mkdir(parents=True)
    q = lambda s: json.dumps(s, ensure_ascii=False)
    meta = dict(fonte="Amostra sintética de aceite A7", aprovacao="Somente teste; sem aprovação financeira real",
                aprovador="Teste", aprovado_em="NAO_INFORMADO", natureza="INTERPRETACAO_TECNICA")
    config = '''versao=1
fonte="Amostra sintética de aceite A7"
[entrada]
aba_taxas="Contratos"
[modalidades]
PIX="pix"
"CREDITO PARCELADO"="parcelado"
'''
    for tabela, r in [
        ('isencoes', dict(id='ISENCAO_TESTE', cnpj='22222222000122', inicio='@2032-10-01', fim='@2032-10-31')),
        ('devolucoes_checkout', dict(id='DEVOLUCAO_TESTE', cnpj='22222222000122', isencao='ISENCAO_TESTE', politica_estornos='VENDA_ORIGINAL_ISENTA')),
        ('contratos_bloqueados', dict(id='SEM_CONTRATO_TESTE', cnpj='33333333000133', motivo='Contrato ainda aguarda o Comercial')),
        ('decisoes_pendentes', dict(id='D-10', descricao='Natureza do ajuste aguarda confirmação financeira', tratamento_atual='Teste da decisão do time, participação de 50% aplicada uma vez'))]:
        config += f'\n[[{tabela}]]\n' + '\n'.join(f'{k} = {v[1:] if isinstance(v,str) and v.startswith("@") else q(v)}' for k,v in dict(r,**meta).items())+'\n'
    for tabela, r in [('remessas', dict(id='REMESSA_TESTE', participacao_fracao='0.5', spread_padrao_fracao='0.015', tolerancia_pp='0.01')),
                      ('notas_checkout', dict(id='NOTA_CHECKOUT_TESTE', servico='Intermediação sintética checkout')),
                      ('notas_remessa', dict(id='NOTA_REMESSA_TESTE', servico='Intermediação sintética remessa', tomador_cnpj='55555555000155', tomador='Contraparte sintética'))]:
        config += f'\n[{tabela}]\n'+'\n'.join(f'{k} = {q(v)}' for k,v in dict(r,**meta).items())+'\n'
    for banco, regras in [('aurora', [('CREDITO', '^credito adquirente$', '1.1.03.001'),('DEBITO','^repasse cliente$','2.1.01.001')]),
                           ('meridiano', [('CREDITO','^recurso cliente$','2.1.02.001'),('DEBITO','^liquidacao cliente$','2.1.02.001')])]:
        for i,(sinal,padrao,conta) in enumerate(regras):
            config += f'\n[[classificacao_{banco}]]\n'+'\n'.join(f'{k} = {q(v)}' for k,v in dict(id=f'{banco}_{i}', sinal=sinal, padrao=padrao, conta_codigo=conta, **meta).items())+'\n'
    regras = raiz/'regras.toml'; regras.write_text(config, encoding='utf-8')
    wb=Workbook(); ws=wb.active; ws.title='Contratos'
    ws.append(['Lojista','CNPJ','EC Orbe','Débito','Crédito à vista','Pix','Antecipação','Spread remessa','Vigência','Obs'])
    for i,cnpj in enumerate(('11111111000111','22222222000122','33333333000133'),900):
        ws.append([f'Cliente sintético {i}',cnpj,str(i),.01,.02,.01,.021,.015,'01/01/2030',''])
    wb.save(dados/'condicoes_novas.xlsx'); wb.close()
    colunas=['id_transacao','data_movimento','cod_estabelecimento','tipo_lancamento','modalidade','qtd_parcelas','valor_bruto','id_referencia','data_venda_original','taxa_aplicada_pct','valor_taxa']
    def orbe(nome, linhas):
        with (dados/nome).open('w',encoding='utf-8',newline='') as f:
            w=csv.writer(f,delimiter=';'); w.writerow(colunas); w.writerows(linhas)
    orbe('primeira_origem.csv',[
        ['VENDA-A','02/10/2032','900','VENDA','PIX','1','1000,00','','','1,1','11,00'],
        ['VENDA-ISENTA','02/10/2032','901','VENDA','PIX','1','100,00','','','2','2,00'],
        ['ANTEC-A','02/10/2032','900','ANTECIPACAO','','','200,00','','02/10/2032','2,1','4,20'],
        ['COMERCIAL','02/10/2032','902','VENDA','PIX','1','100,00','','','1,5','1,50']])
    orbe('chegada_tardia.csv',[
        ['ESTORNO-ISENTO','03/10/2032','901','ESTORNO','PIX','1','-50,00','VENDA-ISENTA','02/10/2032','2','-1,00'],
        ['ESTORNO-ANTERIOR','03/10/2032','901','ESTORNO','PIX','1','-10,00','HISTORICO-SETEMBRO','20/09/2032','1','-0,10'],
        ['REVERSAO-A','03/10/2032','900','REVERSAO ANTECIPACAO','','','-200,00','ANTEC-A','02/10/2032','2,1','-4,20']])
    wb=Workbook();ws=wb.active;ws.title='Remessas novas';ws.append(['Competência: 10/2032'])
    ws.append(['Nº Operação','Data Contratação','Data Liquidação','CNPJ Cliente','Cliente','Moeda','Valor ME','Taxa Cliente','PTAX Ref.','Valor R$','Spread R$','Spread %','Spread Estimado R$ (D-1)','IOF R$','Status Compliance','Observação'])
    regular=['REGULAR-NOVA',datetime(2032,10,2),datetime(2032,11,1),'11111111000111','Cliente sintético 900','USD',100,5.075,5,507.5,10.01,.015,99,0,'APROVADO','']
    ws.append(regular);ws.append(regular[:-1]+['reprocessada'])
    ajuste=regular.copy();ajuste[0]='AJUSTE-NOVO';ajuste[1]=datetime(2032,10,3);ajuste[10]=-2.02;ajuste[-1]='AJUSTE REF REGULAR-NOVA'
    # O contrato do parceiro usa referência PC-… para reconhecer ajuste (RN-37).
    regular_id='PC-OUTRO-MES';ws.cell(3,1,regular_id);ws.cell(4,1,regular_id);ajuste[-1]='AJUSTE REF '+regular_id
    ws.append(ajuste)
    negativo=regular.copy();negativo[0]='NEGATIVO-NOVO';negativo[10]=-2;negativo[11]=-.01;negativo[-1]='Trava anterior'
    ws.append(negativo)
    memoria=BytesIO();wb.save(memoria);wb.close()
    eml=EmailMessage();eml['Subject']='Amostra sintética de outro mês';eml.set_content('Dados locais para teste.')
    eml.add_attachment(memoria.getvalue(),maintype='application',subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',filename='anexo_com_nome_novo.xlsx')
    (dados/'mensagem_nova.eml').write_bytes(eml.as_bytes())
    (dados/'plano_novo.csv').write_text('COD_CONTA;DESCRICAO;TIPO;NATUREZA;ACEITA_LANCTO;COD_SUPERIOR;DT_CRIACAO\n'
        '1.1.03.001;Recebíveis;ANALITICA;D;S;;01/01/2030\n2.1.01.001;Repasses;ANALITICA;C;S;;01/01/2030\n2.1.02.001;Recursos remessa;ANALITICA;C;S;;01/01/2030\n',encoding='latin-1')
    (dados/'banco_csv_novo.csv').write_text('Banco Aurora S.A.;;;;\n;;;;\n;;;;\n;;;;\n;;;;\n;;;;\nData;Lançamento;Documento;Valor (R$);Saldo (R$)\n30/09/2032;SALDO ANTERIOR;;;0,00\n02/10/2032;CREDITO ADQUIRENTE;;999,99;999,99\n03/10/2032;REPASSE CLIENTE;;-123,45;876,54\n',encoding='utf-8')
    s='OFXHEADER:100\nDATA:OFXSGML\nVERSION:102\nENCODING:USASCII\nCHARSET:1252\n\n<OFX>\n<BANKMSGSRSV1>\n<STMTTRNRS>\n<STATUS>\n<CODE>0\n</STATUS>\n<STMTRS>\n<CURDEF>BRL\n<BANKACCTFROM>\n<BANKID>9999\n<ACCTID>OUTRA-CONTA\n<ACCTTYPE>CHECKING\n</BANKACCTFROM>\n<BANKTRANLIST>\n<DTSTART>20321001000000[-3:BRT]\n<DTEND>20321031235959[-3:BRT]\n'
    for ref,tipo,valor,memo in [('FIT-NOVO-A','CREDIT','111.11','RECURSO CLIENTE'),('FIT-NOVO-B','DEBIT','-22.22','LIQUIDACAO CLIENTE')]:
        s+=f'<STMTTRN>\n<TRNTYPE>{tipo}\n<DTPOSTED>20321002120000[-3:BRT]\n<TRNAMT>{valor}\n<FITID>{ref}\n<MEMO>{memo}\n</STMTTRN>\n'
    s+='</BANKTRANLIST>\n</STMTRS>\n</STMTTRNRS>\n</BANKMSGSRSV1>\n</OFX>\n'
    (dados/'banco_ofx_novo.ofx').write_text(s,encoding='cp1252')
    return dados,regras
