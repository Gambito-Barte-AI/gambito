"""A2: taxas assinadas, concessão separada e isolamento do mínimo."""

import csv
from decimal import Decimal
from pathlib import Path
import unittest
from unittest.mock import patch

import test_aceite as amostras
from validar_saida import ErroAceite, verificar
from tarvo.execucao import executar
from tarvo.exportacao import publicar


class A2Test(unittest.TestCase):
    setUp = amostras.AceiteTest.setUp
    vendas = amostras.AceiteTest.vendas

    def orbe(self, linhas):
        colunas = ["id_transacao", "data_movimento", "cod_estabelecimento", "tipo_lancamento",
                   "modalidade", "qtd_parcelas", "valor_bruto", "id_referencia", "data_venda_original",
                   "taxa_aplicada_pct", "valor_taxa"]
        with (self.dados / "orbe.csv").open("w", encoding="utf-8", newline="") as arquivo:
            escritor = csv.DictWriter(arquivo, fieldnames=colunas, delimiter=";")
            escritor.writeheader()
            for l in linhas:
                escritor.writerow(dict(dict.fromkeys(colunas, ""), id_transacao="A", data_movimento="02/04/2032",
                    cod_estabelecimento="42", tipo_lancamento="VENDA", modalidade="PIX", qtd_parcelas="1",
                    valor_bruto="100,00", taxa_aplicada_pct="0,5", valor_taxa="0,50") | l)

    def concessao(self):
        with self.regras.open("a", encoding="utf-8") as f:
            for tabela, corpo in (("isencoes", 'inicio=2032-04-02\nfim=2032-04-30'),
                                  ("devolucoes_checkout", 'isencao="ISENCAO"\npolitica_estornos="VENDA_ORIGINAL_ISENTA"')):
                f.write(f'\n[[{tabela}]]\nid="{"ISENCAO" if tabela == "isencoes" else "DEVOLUCAO"}"\n'
                        'cnpj="01234567800012"\nfonte="Teste sintético"\naprovacao="Amostra"\n'
                        'aprovador="Teste"\naprovado_em="2032-04-01"\nnatureza="FINANCEIRA"\n' + corpo + '\n')

    def publicar(self):
        r = executar(self.dados, self.regras)
        m = publicar(r, self.saida, "alvo")
        verificar(self.dados, self.saida, self.regras, "COMPLETO", "PROVISORIO")
        return r, m

    def test_mesma_taxa_com_truncamento_nao_gera_divergencia(self):
        self.orbe([dict(valor_bruto="1,99", valor_taxa="0,00")])
        r, _ = self.publicar()
        self.assertEqual(r.partes["DIVERGENCIAS_CHECKOUT"].linhas["divergencias_taxa.csv"], [])

    def test_taxas_maiores_menores_estorno_negativo_e_liquido(self):
        self.orbe([dict(taxa_aplicada_pct="1,0"),
                   dict(id_transacao="B", taxa_aplicada_pct="0,4"),
                   dict(id_transacao="E", tipo_lancamento="ESTORNO", valor_bruto="-1,00",
                        taxa_aplicada_pct="1,0", id_referencia="A", data_venda_original="02/04/2032")])
        r, _ = self.publicar()
        parte = r.partes["DIVERGENCIAS_CHECKOUT"]
        self.assertEqual([l["diferenca_brl"] for l in parte.linhas["divergencias_taxa.csv"]], ["0.50", "-0.10", "-0.01"])
        self.assertEqual(parte.linhas["pendencias.csv"][0]["valor_brl"], "0.39")
        self.assertEqual(parte.linhas["pendencias.csv"][0]["ref_origem"], "A")
        self.assertEqual(parte.rastro[2]["ref_evento_original"], "A")

    def test_taxa_diferente_com_valor_arredondado_zero_continua_divergencia(self):
        self.orbe([dict(valor_bruto="0,01", taxa_aplicada_pct="0,6")])
        r, _ = self.publicar()
        self.assertEqual(r.partes["DIVERGENCIAS_CHECKOUT"].linhas["divergencias_taxa.csv"][0]["diferenca_brl"], "0.00")

    def test_vigencia_do_contrato_e_estorno_pela_data_original(self):
        from openpyxl import load_workbook
        livro = load_workbook(self.dados / "taxas.xlsx")
        livro.active.cell(2, 7, "até 02/04/2032")
        livro.active.append(["Loja de teste", "01234567800012", "42", .01, .02, .0025, "03/04/2032", ""])
        livro.save(self.dados / "taxas.xlsx")
        livro.close()
        self.orbe([dict(taxa_aplicada_pct="1,0"),
                   dict(id_transacao="B", data_movimento="03/04/2032", taxa_aplicada_pct="1,0"),
                   dict(id_transacao="E", data_movimento="04/04/2032", tipo_lancamento="ESTORNO",
                        valor_bruto="-100,00", taxa_aplicada_pct="1,0", id_referencia="A", data_venda_original="02/04/2032")])
        r, _ = self.publicar()
        linhas = r.partes["DIVERGENCIAS_CHECKOUT"].linhas["divergencias_taxa.csv"]
        self.assertEqual([l["diferenca_brl"] for l in linhas], ["0.50", "0.75", "-0.50"])

    def test_divergencia_do_estorno_usa_sua_propria_faixa(self):
        from openpyxl import load_workbook
        livro = load_workbook(self.dados / "taxas.xlsx")
        for coluna, titulo, taxa in ((9, "Cred. 2-6x", .03), (10, "Cred. 7-12x", .04)):
            livro.active.cell(1, coluna, titulo)
            livro.active.cell(2, coluna, taxa)
        livro.save(self.dados / "taxas.xlsx")
        livro.close()
        self.regras.write_text(self.regras.read_text(encoding="utf-8").replace('PIX = "pix"',
            'PIX = "pix"\n"CREDITO PARCELADO" = "parcelado"'), encoding="utf-8")
        self.orbe([dict(modalidade="CREDITO PARCELADO", qtd_parcelas="2", taxa_aplicada_pct="3,0"),
                   dict(id_transacao="E", modalidade="CREDITO PARCELADO", qtd_parcelas="7", tipo_lancamento="ESTORNO",
                        valor_bruto="-100,00", taxa_aplicada_pct="5,0", id_referencia="A", data_venda_original="02/04/2032")])
        r, _ = self.publicar()
        p = r.partes["DIVERGENCIAS_CHECKOUT"]
        self.assertEqual(len(p.linhas["divergencias_taxa.csv"]), 1)
        self.assertEqual(p.linhas["divergencias_taxa.csv"][0]["diferenca_brl"], "-1.00")
        self.assertEqual(p.rastro[0]["faixa"], "parcelado_7_12")

    def test_casa_estorno_elegivel_abate_e_anterior_a_isencao_fica_fora(self):
        self.concessao()
        # Antecipação usa sua própria taxa; adicionar a coluna contratual da amostra.
        from openpyxl import load_workbook
        livro = load_workbook(self.dados / "taxas.xlsx")
        livro.active.cell(1, 9, "Antecipação")
        livro.active.cell(2, 9, .01)
        livro.save(self.dados / "taxas.xlsx")
        livro.close()
        self.orbe([dict(id_transacao="ANTES", data_movimento="01/04/2032"), dict(),
                   dict(id_transacao="ULTIMA", data_movimento="30/04/2032"),
                   dict(id_transacao="ANT", tipo_lancamento="ANTECIPACAO", taxa_aplicada_pct="1,0", valor_taxa="1,23"),
                   dict(id_transacao="EA", tipo_lancamento="ESTORNO", valor_bruto="-100,00", valor_taxa="-0,50",
                        id_referencia="ANTES", data_venda_original="01/04/2032"),
                   dict(id_transacao="E", tipo_lancamento="ESTORNO", valor_bruto="-100,00", valor_taxa="-0,50",
                        id_referencia="A", data_venda_original="02/04/2032")])
        r, _ = self.publicar()
        parte = r.partes["DEVOLUCAO_CASA"]
        self.assertEqual(parte.situacao, "COMPLETO")
        self.assertEqual(parte.linhas["pendencias.csv"][0]["valor_brl"], "1.73")
        self.assertEqual({l["ref_origem"] for l in parte.rastro}, {"A", "ANT", "ULTIMA", "E"})
        estorno = next(l for l in parte.rastro if l["ref_origem"] == "E")
        self.assertEqual((estorno["valor"], estorno["ref_evento_original"], estorno["arquivo_evento_original"]),
                         ("-0.50", "A", "orbe.csv"))
        self.assertFalse(r.partes["DIVERGENCIAS_CHECKOUT"].rastro)
        self.assertEqual(next(c.valor for c in r.contribuicoes if c.origem.ref == "EA"), Decimal("-.50"))

    def test_estorno_apos_fim_isencao_usa_venda_original_isenta_em_outro_csv(self):
        self.concessao()
        self.regras.write_text(self.regras.read_text(encoding="utf-8").replace("fim=2032-04-30", "fim=2032-04-03"), encoding="utf-8")
        self.orbe([dict()])
        with (self.dados / "orbe.csv").open(encoding="utf-8", newline="") as f:
            colunas = next(csv.reader(f, delimiter=";"))
        with (self.dados / "ajuste_sem_nome_fixo.csv").open("w", encoding="utf-8", newline="") as f:
            escritor = csv.DictWriter(f, fieldnames=colunas, delimiter=";")
            escritor.writeheader()
            escritor.writerow(dict(id_transacao="E", data_movimento="20/04/2032", cod_estabelecimento="42",
                tipo_lancamento="ESTORNO", modalidade="PIX", qtd_parcelas="1", valor_bruto="-50,00",
                id_referencia="A", data_venda_original="02/04/2032", taxa_aplicada_pct="0,5", valor_taxa="-0,25"))
        r, _ = self.publicar()
        p = r.partes["DEVOLUCAO_CASA"]
        self.assertEqual(p.linhas["pendencias.csv"][0]["valor_brl"], "0.25")
        estorno = next(l for l in p.rastro if l["ref_origem"] == "E")
        self.assertEqual(estorno["arquivo_origem"], "ajuste_sem_nome_fixo.csv")
        self.assertEqual(estorno["arquivo_evento_original"], "orbe.csv")
        self.assertEqual(estorno["data_base_taxa"], "2032-04-02")

    def test_estorno_sem_original_do_mes_nao_e_abate_estimado(self):
        self.concessao()
        self.orbe([dict(), dict(id_transacao="E", tipo_lancamento="ESTORNO", valor_bruto="-100,00",
                               valor_taxa="-0,50", id_referencia="AUSENTE", data_venda_original="02/04/2032")])
        r = executar(self.dados, self.regras)
        m = publicar(r, self.saida, "alvo")
        p = r.partes["DEVOLUCAO_CASA"]
        self.assertEqual(p.situacao, "PROVISORIO")
        self.assertEqual(p.linhas["pendencias.csv"][0]["valor_brl"], "0.50")
        self.assertEqual([l["ref_origem"] for l in p.rastro], ["A"])
        self.assertTrue(any(p.origem.ref == "E" for p in r.pendencias))
        verificar(self.dados, self.saida, self.regras, "PROVISORIO")

    def test_estorno_isento_sem_cobrado_fica_pendente_sem_inventar_abate(self):
        self.concessao()
        self.orbe([dict(), dict(id_transacao="E", tipo_lancamento="ESTORNO", valor_bruto="-100,00",
                               valor_taxa="inválido", id_referencia="A", data_venda_original="02/04/2032")])
        r, _ = self.publicar()
        p = r.partes["DEVOLUCAO_CASA"]
        self.assertEqual(p.situacao, "PROVISORIO")
        self.assertEqual(p.cobertura["pendentes"], 1)
        revisao = next(l for l in p.linhas["pendencias.csv"] if l["id_saida"].startswith("REVISAO_"))
        self.assertEqual(revisao["valor_brl"], "")

    def test_auditor_rejeita_omissao_do_estorno_isento_com_soma_reconciliada(self):
        self.concessao()
        self.orbe([dict(), dict(id_transacao="E", tipo_lancamento="ESTORNO", valor_bruto="-50,00",
                               valor_taxa="-0,25", id_referencia="A", data_venda_original="02/04/2032")])
        r = executar(self.dados, self.regras)
        p = r.partes["DEVOLUCAO_CASA"]
        p.rastro = [l for l in p.rastro if l["ref_origem"] != "E"]
        p.linhas["pendencias.csv"][0]["valor_brl"] = "0.50"
        publicar(r, self.saida, "alvo")
        with self.assertRaises(ErroAceite):
            verificar(self.dados, self.saida, self.regras)

    def test_taxa_aplicada_invalida_so_afeta_divergencia(self):
        self.orbe([dict(taxa_aplicada_pct="inválida")])
        r, _ = self.publicar()
        self.assertEqual(r.situacao, "COMPLETO")
        self.assertEqual(r.partes["DIVERGENCIAS_CHECKOUT"].situacao, "PROVISORIO")
        self.assertEqual(r.partes["DEVOLUCAO_CASA"].situacao, "COMPLETO")

    def test_cobrado_invalido_so_afeta_devolucao_sem_inventar_zero(self):
        self.concessao()
        self.orbe([dict(valor_taxa="inválido")])
        r, _ = self.publicar()
        p = r.partes["DEVOLUCAO_CASA"]
        self.assertEqual(p.situacao, "PROVISORIO")
        self.assertEqual(p.linhas["pendencias.csv"][0]["valor_brl"], "")
        self.assertFalse(p.rastro)
        self.assertEqual(p.cobertura["pendentes"], 1)
        self.assertEqual(r.partes["DIVERGENCIAS_CHECKOUT"].situacao, "COMPLETO")

    def test_excecao_de_uma_analise_nao_bloqueia_irma_ou_minimo(self):
        self.orbe([dict()])
        with patch("tarvo.divergencias.calcular_devolucao", side_effect=ValueError("Falha sintética")):
            r, m = self.publicar()
        self.assertEqual(m["partes"]["DEVOLUCAO_CASA"]["situacao"], "BLOQUEADO")
        self.assertEqual(r.partes["DIVERGENCIAS_CHECKOUT"].situacao, "COMPLETO")

    def test_rastro_divergencia_corrompido_preserva_outros(self):
        self.orbe([dict(taxa_aplicada_pct="1,0")])
        r = executar(self.dados, self.regras)
        r.partes["DIVERGENCIAS_CHECKOUT"].rastro[0]["valor"] = "9.99"
        m = publicar(r, self.saida, "alvo")
        self.assertEqual(m["partes"]["DIVERGENCIAS_CHECKOUT"]["situacao"], "BLOQUEADO")
        self.assertEqual(m["partes"]["CHECKOUT"]["situacao"], "COMPLETO")
        self.assertEqual(m["partes"]["DEVOLUCAO_CASA"]["situacao"], "COMPLETO")
        verificar(self.dados, self.saida, self.regras, "COMPLETO")

    def test_comissao_corrompida_preserva_divergencias(self):
        self.orbe([dict(taxa_aplicada_pct="1,0")])
        r = executar(self.dados, self.regras)
        r.rastro[0]["valor"] = "9.99"
        m = publicar(r, self.saida, "alvo")
        self.assertEqual(m["partes"]["CHECKOUT"]["situacao"], "BLOQUEADO")
        self.assertEqual(m["partes"]["DIVERGENCIAS_CHECKOUT"]["situacao"], "COMPLETO")
        verificar(self.dados, self.saida, self.regras, "BLOQUEADO")

    def test_auditor_rejeita_formula_alterada_mesmo_com_rastro_reconciliado(self):
        self.orbe([dict(taxa_aplicada_pct="1,0")])
        r = executar(self.dados, self.regras)
        p = r.partes["DIVERGENCIAS_CHECKOUT"]
        p.linhas["divergencias_taxa.csv"][0]["diferenca_brl"] = "9.99"
        p.linhas["pendencias.csv"][0]["valor_brl"] = "9.99"
        for l in p.rastro:
            l["valor"] = "9.99"
        publicar(r, self.saida, "alvo")
        with self.assertRaises(ErroAceite):
            verificar(self.dados, self.saida, self.regras)

    def test_politica_aprovada_real_usa_isencao_original_e_preserva_d10(self):
        from tarvo.regras import carregar
        config = carregar(Path(__file__).resolve().parents[1] / "regras.toml")
        self.assertEqual(config["devolucoes_checkout"][0]["politica_estornos"], "VENDA_ORIGINAL_ISENTA")
        self.assertEqual([d["id"] for d in config["decisoes_pendentes"]], ["D-10"])

