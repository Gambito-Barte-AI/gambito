"""Aceite do comando instalado e auditoria independente, fora da raiz do projeto."""

import csv
from hashlib import sha256
import json
import os
from pathlib import Path
import subprocess
import sys
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from validar_saida import ErroAceite, verificar


class AceiteTest(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.raiz = Path(self.tmp.name)
        self.dados = self.raiz / "dados"
        self.dados.mkdir()
        self.saida = self.raiz / "saida"
        self.regras = self.raiz / "regras.toml"
        self.regras.write_text('''versao = 1
fonte = "Amostra sintética de aceite"
[entrada]
aba_taxas = "Taxas"
[modalidades]
PIX = "pix"
[[decisoes_pendentes]]
id = "D-07"
descricao = "Decisão não ratificada"
tratamento_atual = "Amostra de verificação da preservação das decisões"
''', encoding="utf-8")
        livro = Workbook()
        aba = livro.active
        aba.title = "Taxas"
        aba.append(["Lojista", "CNPJ", "EC Orbe", "Débito", "Crédito à vista", "Pix", "Vigência", "Obs"])
        aba.append(["Loja de teste", "01234567800012", "42", .01, .02, .005, "01/01/2030", ""])
        livro.save(self.dados / "taxas.xlsx")
        livro.close()
        self.vendas()

    def vendas(self, extras=(), data="02/04/2032"):
        with (self.dados / "orbe.csv").open("w", encoding="utf-8", newline="") as arquivo:
            escritor = csv.writer(arquivo, delimiter=";")
            escritor.writerow(["id_transacao", "data_movimento", "cod_estabelecimento", "tipo_lancamento",
                               "modalidade", "qtd_parcelas", "valor_bruto", "id_referencia", "data_venda_original"])
            escritor.writerow(["A", data, "42", "VENDA", "PIX", "1", "1,00", "", ""])
            for ref, modalidade, ec, parcelas in extras:
                escritor.writerow([ref, data, ec, "VENDA", modalidade, parcelas, "1,00", "", ""])

    def executar(self, codigo=0):
        ambiente = os.environ.copy()
        for chave in ("PYTHONPATH", "PYTHONHOME", "ANTHROPIC_API_KEY", "OPENAI_API_KEY"):
            ambiente.pop(chave, None)
        ambiente["PYTHONNOUSERSITE"] = "1"
        ambiente["PYTHONUTF8"] = "1"
        processo = subprocess.run(
            [sys.executable, "-m", "tarvo", "--dados", str(self.dados), "--saida", str(self.saida),
             "--regras", str(self.regras)], cwd=self.raiz, env=ambiente,
            capture_output=True, text=True, encoding="utf-8", timeout=60,
        )
        self.assertEqual(processo.returncode, codigo, processo.stdout + processo.stderr)
        return processo

    def auditar(self, esperar="COMPLETO"):
        return verificar(self.dados, self.saida, self.regras, esperar)

    def arquivo_orbe(self, nome, linhas):
        colunas = ["id_transacao", "data_movimento", "cod_estabelecimento", "tipo_lancamento",
                   "modalidade", "qtd_parcelas", "valor_bruto", "id_referencia", "data_venda_original", "valor_taxa"]
        with (self.dados / nome).open("w", encoding="utf-8", newline="") as arquivo:
            escritor = csv.DictWriter(arquivo, fieldnames=colunas, delimiter=";")
            escritor.writeheader()
            for campos in linhas:
                linha = dict.fromkeys(colunas, "")
                linha.update(id_transacao="A", data_movimento="02/04/2032", cod_estabelecimento="42",
                             tipo_lancamento="VENDA", modalidade="PIX", qtd_parcelas="1", valor_bruto="1,00")
                linha.update(campos)
                escritor.writerow(linha)

    def test_ajuste_em_outro_arquivo_usa_original_e_preserva_origem(self):
        self.arquivo_orbe("qualquer_nome.csv", [dict(id_transacao="E", data_movimento="03/04/2032",
            tipo_lancamento="ESTORNO", valor_bruto="-1,00", id_referencia="A", data_venda_original="02/04/2032")])
        self.executar()
        self.assertEqual(self.auditar()["total_brl"], "0.00")
        with (self.saida / "rastro.csv").open(encoding="utf-8") as arquivo:
            rastros = list(csv.DictReader(arquivo, delimiter=";"))
        estorno = next(r for r in rastros if r["ref_origem"] == "E")
        self.assertEqual((estorno["arquivo_origem"], estorno["arquivo_evento_original"]),
                         ("qualquer_nome.csv", "orbe.csv"))
        self.assertEqual(estorno["data_base_taxa"], "2032-04-02")

    def test_reversao_em_outro_arquivo_mesmo_antes_do_original(self):
        from openpyxl import load_workbook
        livro = load_workbook(self.dados / "taxas.xlsx")
        livro["Taxas"].cell(1, 9, "Antecipação")
        livro["Taxas"].cell(2, 9, .021)
        livro.save(self.dados / "taxas.xlsx")
        livro.close()
        self.arquivo_orbe("orbe.csv", [dict(tipo_lancamento="ANTECIPACAO", valor_bruto="100,00")])
        self.arquivo_orbe("antes.csv", [dict(id_transacao="R", data_movimento="03/04/2032",
            tipo_lancamento="REVERSAO ANTECIPACAO", valor_bruto="-100,00", id_referencia="A", data_venda_original="02/04/2032")])
        self.executar()
        self.assertEqual(self.auditar()["total_brl"], "0.00")

    def test_duplicata_identica_uma_comissao_e_dois_destinos_fisicos(self):
        (self.dados / "copia.csv").write_bytes((self.dados / "orbe.csv").read_bytes())
        self.executar(2)
        aceite = self.auditar("PROVISORIO")
        self.assertEqual(aceite["total_brl"], "0.01")
        self.assertEqual(aceite["cobertura"], dict(linhas_lidas=2, calculadas=1, pendentes=1, bloqueadas=0))

    def test_auditor_detecta_csv_orbe_omitido(self):
        self.executar()
        self.arquivo_orbe("recebido_depois.csv", [dict(id_transacao="B")])
        with self.assertRaisesRegex(ErroAceite, "omite.*CSV Orbe"):
            self.auditar()

    def test_coluna_opcional_com_campo_faltante_vira_pendencia(self):
        self.arquivo_orbe("ajuste.csv", [dict(id_transacao="B")])
        caminho = self.dados / "ajuste.csv"
        texto = caminho.read_text(encoding="utf-8")
        caminho.write_text(texto.rstrip("\n").removesuffix(";") + "\n", encoding="utf-8")
        self.executar(2)
        self.assertEqual(self.auditar("PROVISORIO")["cobertura"]["pendentes"], 1)

    def test_duplicata_com_ocorrencia_invalida_nao_libera_comissao(self):
        self.arquivo_orbe("invalido.csv", [dict(valor_bruto="NaN")])
        self.executar(2)
        aceite = self.auditar("PROVISORIO")
        self.assertEqual(aceite["cobertura"]["calculadas"], 0)
        self.assertEqual(aceite["cobertura"]["pendentes"], 2)

    def test_ids_conflitantes_nao_escolhem_valor_nem_liberam_referencia(self):
        self.arquivo_orbe("conflito.csv", [dict(valor_bruto="2,00")])
        self.arquivo_orbe("ajuste.csv", [dict(id_transacao="E", tipo_lancamento="ESTORNO",
            valor_bruto="-1,00", id_referencia="A", data_venda_original="02/04/2032")])
        self.executar(2)
        aceite = self.auditar("PROVISORIO")
        self.assertEqual(aceite["cobertura"], dict(linhas_lidas=3, calculadas=0, pendentes=3, bloqueadas=0))

    def test_orbe_valida_e_adicional_com_layout_quebrado_fecha_provisorio(self):
        (self.dados / "ruim.csv").write_text("id_transacao;outra\nB;1\n", encoding="utf-8")
        self.executar(2)
        self.assertEqual(self.auditar("PROVISORIO")["total_brl"], "0.01")

    def test_mudanca_na_coluna_id_nao_descarta_arquivo_silenciosamente(self):
        self.arquivo_orbe("layout_mudou.csv", [dict(id_transacao="B")])
        caminho = self.dados / "layout_mudou.csv"
        caminho.write_text(caminho.read_text(encoding="utf-8").replace("id_transacao;", "novo_id;"), encoding="utf-8")
        self.executar(2)
        self.assertEqual(self.auditar("PROVISORIO")["total_brl"], "0.01")

    def test_referencia_do_mes_ausente_entra_com_pendencia_derivada(self):
        self.arquivo_orbe("ajuste.csv", [dict(id_transacao="E", tipo_lancamento="ESTORNO",
            valor_bruto="-1,00", id_referencia="NAO_EXISTE", data_venda_original="02/04/2032")])
        self.executar(2)
        aceite = self.auditar("PROVISORIO")
        self.assertEqual(aceite["cobertura"]["calculadas"], 2)
        self.assertEqual(aceite["cobertura"]["pendentes"], 0)
        with (self.saida / "pendencias.csv").open(encoding="utf-8", newline="") as arquivo:
            linhas = list(csv.DictReader(arquivo, delimiter=";"))
        self.assertEqual(len([l for l in linhas if l["id_saida"].startswith("AJUSTE_ORFAO_ORBE:")]), 1)

    def test_escolha_explicita_de_competencia_e_mes_invalido(self):
        from tarvo.execucao import executar
        from tarvo.exportacao import publicar
        self.arquivo_orbe("outro_mes.csv", [dict(id_transacao="B", data_movimento="01/05/2032")])
        self.assertEqual(executar(self.dados, self.regras).situacao, "BLOQUEADO")
        resultado = executar(self.dados, self.regras, "2032-05")
        self.assertEqual(resultado.situacao, "COMPLETO")
        self.assertEqual(resultado.ingestao["fora_competencia"], 1)
        publicar(resultado, self.saida)
        self.assertEqual(self.auditar()["competencia"], "2032-05")
        self.assertEqual(executar(self.dados, self.regras, "2032-13").situacao, "BLOQUEADO")

    def test_pendencia_comercial_preserva_cobrado_sem_tornar_provisorio(self):
        with self.regras.open("a", encoding="utf-8") as arquivo:
            arquivo.write('''
[[contratos_bloqueados]]
id = "SEM_CONTRATO"
cnpj = "01234567800012"
motivo = "Comercial deve confirmar contrato; não faturar"
fonte = "Amostra sintética"
aprovacao = "Teste"
aprovador = "Teste sintético"
aprovado_em = "2032-01-01"
natureza = "FINANCEIRA"
''')
        self.arquivo_orbe("orbe.csv", [dict(valor_taxa="12,34")])
        self.executar()
        self.auditar()
        manifesto = json.loads((self.saida / "manifesto.json").read_text(encoding="utf-8"))
        pendencia = manifesto["pendencias"][0]
        self.assertEqual(pendencia["valor_orbe_brl"], "12.34")
        self.assertEqual(pendencia["valor_bruto_brl"], "1.00")
        self.assertEqual(pendencia["categoria"], "NEGOCIO")
        self.assertTrue(pendencia["proximo_passo"])
        self.assertEqual(manifesto["cobertura"]["calculadas"], 0)
        with (self.saida / "pendencias.csv").open(encoding="utf-8") as arquivo:
            linhas = list(csv.DictReader(arquivo, delimiter=";"))
        self.assertEqual(len(linhas), 1)
        self.assertEqual(linhas[0]["valor_brl"], "12.34")
        self.assertEqual(linhas[0]["natureza_valor"], "COBRANCA_ORBE")
        self.assertTrue(all(linhas[0][k] for k in ("tipo", "ref_origem", "motivo", "quem_decide", "proximo_passo")))
        self.assertEqual(manifesto["resultados"]["pendencias.csv"]["situacao"], "COMPLETO")
        for nome in ("divergencias_taxa.csv", "notas_a_emitir.csv", "lancamentos_classificados.csv"):
            self.assertEqual(manifesto["resultados"][nome]["situacao"], "BLOQUEADO")
            self.assertEqual(len((self.saida / nome).read_text().splitlines()), 1)

    def test_pendencia_sem_valor_nao_vira_zero_nem_rastro_monetario(self):
        self.arquivo_orbe("invalido.csv", [dict(id_transacao="B", valor_bruto="NaN")])
        self.executar(2)
        aceite = self.auditar("PROVISORIO")
        self.assertEqual(aceite["pendencias_csv"], 1)
        self.assertEqual(aceite["rastro_pendencias"], 0)
        with (self.saida / "pendencias.csv").open(encoding="utf-8") as arquivo:
            p = next(csv.DictReader(arquivo, delimiter=";"))
        self.assertEqual(p["valor_brl"], "")
        self.assertEqual(p["natureza_valor"], "NAO_INFORMADO")

    def test_auditor_detecta_valor_adulterado_na_pendencia(self):
        self.vendas([("B", "NOVA", "42", "1")])
        self.executar(2)
        caminho = self.saida / "pendencias.csv"
        caminho.write_bytes(caminho.read_bytes().replace(b";1.00;", b";9.99;"))
        self.regravar_hash(caminho.name)
        with self.assertRaisesRegex(ErroAceite, "Valor da pendência CSV|Soma do rastro"):
            self.auditar("PROVISORIO")

    def test_auditor_detecta_rastro_de_pendencia_adulterado(self):
        self.vendas([("B", "NOVA", "42", "1")])
        self.executar(2)
        caminho = self.saida / "rastro.csv"
        caminho.write_bytes(caminho.read_bytes().replace(b";1.00;", b";9.99;"))
        self.regravar_hash(caminho.name)
        with self.assertRaisesRegex(ErroAceite, "Rastro da pendência|Soma do rastro"):
            self.auditar("PROVISORIO")

    def test_referencias_documentais_e_decisoes_resolvidas(self):
        import re
        import tomllib
        raiz = Path(__file__).resolve().parents[1]
        config = tomllib.loads((raiz / "regras.toml").read_text(encoding="utf-8"))
        documento = raiz / "docs/documentação-geral.md"
        self.assertTrue(documento.is_file())
        texto = documento.read_text(encoding="utf-8")
        self.assertFalse({"D-07", "D-08"} & {d["id"] for d in config.get("decisoes_pendentes", [])})
        for tabela in ("vinculos", "contratos_bloqueados", "isencoes", "observacoes_esclarecidas", "decisoes_documentadas", "decisoes_pendentes"):
            for regra in config.get(tabela, []):
                self.assertIn("docs/documentação-geral.md", regra["fonte"])
                for referencia in re.findall(r"\b(?:RN|D)-\d+\b", regra["fonte"]):
                    self.assertIn(referencia, texto)

    def test_seis_cabecalhos_seguem_o_contrato_oficial_em_ordem(self):
        import re
        from tarvo.exportacao import ARQUIVOS
        raiz = Path(__file__).resolve().parents[1]
        texto = (raiz / "docs/FORMATO_SAIDA.md").read_text(encoding="utf-8")
        contratos = dict(re.findall(r"^\| `([^`]+\.csv)` \| `([^`]+)` \|$", texto, re.MULTILINE))
        self.assertEqual(set(contratos), set(ARQUIVOS))
        for nome, texto_colunas in contratos.items():
            obrigatorias = [c.strip() for c in texto_colunas.split(";")]
            self.assertEqual(ARQUIVOS[nome][:len(obrigatorias)], obrigatorias)

    def test_relatorio_alvo_existente_do_usuario_nao_e_sobrescrito(self):
        self.saida.mkdir()
        caminho = self.saida / "notas_a_emitir.csv"
        caminho.write_text("arquivo externo", encoding="utf-8")
        self.executar(1)
        self.assertEqual(caminho.read_text(encoding="utf-8"), "arquivo externo")
        self.assertFalse((self.saida / "manifesto.json").exists())

    def test_metadados_de_aprovacao_obrigatorios_sem_inventar_datas(self):
        from tarvo.regras import carregar
        from tarvo.modelos import ErroEntrada
        original = Path(__file__).resolve().parents[1] / "regras.toml"
        texto = original.read_text(encoding="utf-8")
        config = carregar(original)
        self.assertTrue(all(r["aprovado_em"] != "NAO_INFORMADO" for r in config["observacoes_esclarecidas"]))
        # A fonte atual documenta as datas; uma fonte futura sem evidência deve
        # poder declarar a ausência, sem inventar aprovação financeira.
        self.regras.write_text(texto.replace('2026-10-16T20:24:00-03:00', 'NAO_INFORMADO', 1), encoding="utf-8")
        self.assertEqual(carregar(self.regras)["vinculos"][0]["aprovado_em"], "NAO_INFORMADO")
        cabecalho, vinculos = texto.split('[[vinculos]]', 1)
        self.regras.write_text(cabecalho + '[[vinculos]]' + vinculos.replace('aprovador = ', 'aprovador_omitido = ', 1), encoding="utf-8")
        with self.assertRaisesRegex(ErroEntrada, "RN-61"):
            carregar(self.regras)
        self.regras.write_text(texto.replace('2026-10-16T20:24:00-03:00', 'ontem', 1), encoding="utf-8")
        with self.assertRaisesRegex(ErroEntrada, "data de aprovação"):
            carregar(self.regras)

    def regravar_hash(self, nome):
        caminho = self.saida / "manifesto.json"
        m = json.loads(caminho.read_text(encoding="utf-8"))
        m["arquivos"][nome] = sha256((self.saida / nome).read_bytes()).hexdigest()
        caminho.write_text(json.dumps(m), encoding="utf-8")

    def test_execucao_instalada_fora_da_raiz_sem_chaves_e_repeticao(self):
        self.executar()
        antes = {p.name: p.read_bytes() for p in self.saida.iterdir()}
        self.executar()
        self.assertEqual(antes, {p.name: p.read_bytes() for p in self.saida.iterdir()})
        self.assertEqual(self.auditar()["total_brl"], "0.01")

    def test_outro_mes_e_modalidade_desconhecida_sem_perda_de_cobertura(self):
        self.vendas([("B", "NOVA", "42", "1"), ("C", "PIX", "99", "1"),
                     ("D", "PIX", "42", "13")], data="02/05/2032")
        self.executar(2)
        aceite = self.auditar("PROVISORIO")
        self.assertEqual(aceite["competencia"], "2032-05")
        self.assertEqual(aceite["cobertura"], {"linhas_lidas": 4, "calculadas": 1, "pendentes": 3, "bloqueadas": 0})

    def test_fonte_ausente_invalida_saida_anterior(self):
        self.executar()
        (self.dados / "taxas.xlsx").unlink()
        self.executar(1)
        self.auditar("BLOQUEADO")
        self.assertEqual(len((self.saida / "rastro.csv").read_text().splitlines()), 1)
        self.assertEqual(len((self.saida / "comissoes_por_lojista.csv").read_text().splitlines()), 1)

    def test_layout_invalido_publica_apenas_cabecalhos(self):
        (self.dados / "orbe.csv").write_text("id_transacao;outra\nA;1\n", encoding="utf-8")
        self.executar(1)
        self.auditar("BLOQUEADO")
        self.assertEqual(len(list(self.saida.glob("*.csv"))), 6)
        for p in self.saida.glob("*.csv"):
            self.assertEqual(len(p.read_text().splitlines()), 1)

    def test_auditor_detecta_soma_errada_mesmo_com_hash_atualizado(self):
        self.executar()
        caminho = self.saida / "comissoes_por_lojista.csv"
        caminho.write_bytes(caminho.read_bytes().replace(b"0.01", b"9.99"))
        self.regravar_hash(caminho.name)
        with self.assertRaisesRegex(ErroAceite, "Soma do rastro"):
            self.auditar()

    def test_auditor_detecta_referencia_inventada_mesmo_com_hash_atualizado(self):
        self.executar()
        caminho = self.saida / "rastro.csv"
        caminho.write_bytes(caminho.read_bytes().replace(b";A;", b";INVENTADA;"))
        self.regravar_hash(caminho.name)
        with self.assertRaisesRegex(ErroAceite, "destino"):
            self.auditar()

    def test_auditor_detecta_decisao_financeira_omitida(self):
        self.executar()
        caminho = self.saida / "manifesto.json"
        m = json.loads(caminho.read_text(encoding="utf-8"))
        m["decisoes_financeiras_pendentes"] = []
        caminho.write_text(json.dumps(m), encoding="utf-8")
        with self.assertRaisesRegex(ErroAceite, "Decisões financeiras"):
            self.auditar()

    def test_soma_interna_corrompida_bloqueia_publicacao(self):
        from tarvo.execucao import executar
        from tarvo.exportacao import publicar
        from tarvo.validacao import consolidar

        def corromper(resultado):
            consolidar(resultado)
            resultado.rastro[0]["valor"] = "9.99"

        with patch("tarvo.execucao.validacao.consolidar", side_effect=corromper):
            resultado = executar(self.dados, self.regras)
        self.assertEqual(resultado.situacao, "BLOQUEADO")
        publicar(resultado, self.saida)
        self.auditar("BLOQUEADO")
        self.assertEqual(len(list(self.saida.glob("*.csv"))), 6)
        for p in self.saida.glob("*.csv"):
            self.assertEqual(len(p.read_text().splitlines()), 1)

    def test_publicacao_interrompida_nao_deixa_manifesto_valido(self):
        from tarvo.execucao import executar
        from tarvo.exportacao import publicar

        original = os.replace
        chamadas = []

        def falhar_no_segundo_csv(origem, destino):
            chamadas.append(str(destino))
            if len(chamadas) == 3:
                raise OSError("falha de escrita simulada")
            return original(origem, destino)

        with patch("tarvo.exportacao.os.replace", side_effect=falhar_no_segundo_csv):
            with self.assertRaises(OSError):
                publicar(executar(self.dados, self.regras), self.saida)
        self.auditar("BLOQUEADO")


if __name__ == "__main__":
    unittest.main()
